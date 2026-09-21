"""Reusable Presentation Primitives for Rise of Empires Editorial Theme.

Generates editable PowerPoint shapes, shadows, borders, image frames, and typography
matching the exact XML properties of benchmark deck 'The Rise of Empires (3).pptx'.
"""

import io
import math
import logging
import os
from typing import Any, List, Optional, Tuple

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from PIL import Image as PILImage

from app.presentation.themes.rise_of_empires_editorial.tokens import COLORS, SHADOWS, SPACING, FONTS, GEOMETRY
from app.presentation.themes.rise_of_empires_editorial.icons import IconResolver
from app.presentation.themes.rise_of_empires_editorial.typography import (
    BUDGETS, fit_text_to_box, format_kicker_text, TitleOptimizer, TextMeasurementService
)

logger = logging.getLogger(__name__)


def apply_soft_shadow(
    shape: Any,
    blur_rad: int = SHADOWS.blur_rad_emu,
    dist: int = SHADOWS.dist_emu,
    dir_val: int = SHADOWS.dir_deg,
    color: str = SHADOWS.color_hex,
    alpha: int = SHADOWS.alpha_val
):
    """Applies DrawingML outer shadow (16% black, 9pt blur, 3pt distance, 90 deg down)."""
    try:
        spPr = shape._element.spPr
        effectLst = spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        if effectLst is None:
            effectLst = OxmlElement("a:effectLst")
            spPr.append(effectLst)

        outerShdw = OxmlElement("a:outerShdw")
        outerShdw.set("blurRad", str(blur_rad))
        outerShdw.set("dist", str(dist))
        outerShdw.set("dir", str(dir_val))
        outerShdw.set("algn", "tl")
        outerShdw.set("rotWithShape", "0")

        srgbClr = OxmlElement("a:srgbClr")
        srgbClr.set("val", color)

        alpha_el = OxmlElement("a:alpha")
        alpha_el.set("val", str(alpha))
        srgbClr.append(alpha_el)
        outerShdw.append(srgbClr)

        effectLst.append(outerShdw)
    except Exception as e:
        logger.debug("Failed to set outer shadow: %s", e)


def set_shape_border(shape: Any, color: RGBColor, width_pt: float = 0.75):
    """Sets a clean vector outline on a shape."""
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def create_solid_shape(
    slide: Any,
    shape_type: MSO_SHAPE,
    x: float,
    y: float,
    w: float,
    h: float,
    fill_color: RGBColor,
    border_color: Optional[RGBColor] = None,
    border_width_pt: float = 0.75,
    shadow: bool = False,
    corner_radius_pct: Optional[float] = None
) -> Any:
    """Creates a solid shape with clean styling, subtle corner radius, and optional shadow."""
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_color
    if border_color:
        set_shape_border(shp, border_color, border_width_pt)
    else:
        shp.line.fill.background()
    if shadow:
        apply_soft_shadow(shp)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            if corner_radius_pct is not None:
                shp.adjustments[0] = corner_radius_pct
            else:
                # Authentic editorial corner radius: ~0.16 inches subtle curve, preventing bulbous joining
                min_dim = max(0.1, min(w, h))
                shp.adjustments[0] = min(0.06, max(0.022, 0.16 / min_dim))
        except Exception:
            pass
    return shp


def add_text_box(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    margin_left: float = 0.0,
    margin_top: float = 0.0,
    margin_right: float = 0.0,
    margin_bottom: float = 0.0,
    name: Optional[str] = None
) -> Any:
    """Creates a text box with tight zero margins for precision alignment."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margin_left)
    tf.margin_right = Inches(margin_right)
    tf.margin_top = Inches(margin_top)
    tf.margin_bottom = Inches(margin_bottom)
    return tb


class IconBadge:
    """Renders the signature burnt-orange circular badge with centered icon asset."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        diameter: float = 0.66,
        icon_path: Optional[str] = None,
        semantic_hint: Optional[str] = None,
        slide_idx: Optional[int] = None
    ) -> Any:
        # 1. Circular badge shape
        circle = create_solid_shape(
            slide,
            MSO_SHAPE.OVAL,
            x,
            y,
            diameter,
            diameter,
            fill_color=COLORS.rgb_accent,
            border_color=None,
            shadow=False
        )

        # 2. Resolve icon asset
        resolved_icon = icon_path
        if not resolved_icon:
            resolved_icon = IconResolver.resolve_path(
                semantic_hint=semantic_hint,
                slide_idx=slide_idx
            )

        # 3. Add icon image perfectly centered
        if resolved_icon and os.path.exists(resolved_icon):
            ix, iy, iw, ih = IconResolver.calculate_badge_icon_geometry(
                badge_x=x,
                badge_y=y,
                badge_diameter=diameter,
                glyph_ratio=0.52
            )
            slide.shapes.add_picture(resolved_icon, Inches(ix), Inches(iy), Inches(iw), Inches(ih))

        return circle


class NumberBadge:
    """Renders burnt-orange circular badge with centered two-digit number."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        number_str: str,
        diameter: float = 0.66,
        font_size_pt: float = 14.0
    ) -> Any:
        circle = create_solid_shape(
            slide,
            MSO_SHAPE.OVAL,
            x,
            y,
            diameter,
            diameter,
            fill_color=COLORS.rgb_accent,
            border_color=None,
            shadow=False
        )
        tf = circle.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        p = tf.paragraphs[0]
        p.text = str(number_str).zfill(2)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONTS.body
        p.font.size = Pt(font_size_pt)
        p.font.bold = True
        p.font.color.rgb = COLORS.rgb_white
        return circle


class HistoricalSlideHeader:
    """Renders the standard header: Icon Badge, Antique Gold Kicker, and Cambria Title."""

    @classmethod
    def render(
        cls,
        slide: Any,
        kicker: str,
        title: str,
        slide_idx: Optional[int] = None,
        semantic_hint: Optional[str] = None,
        icon_path: Optional[str] = None,
        is_dark: bool = False
    ) -> Tuple[Any, Any, Any]:
        badge_y = GEOMETRY.dark_badge_y if is_dark else GEOMETRY.header_badge_y
        kicker_y = GEOMETRY.dark_kicker_y if is_dark else GEOMETRY.header_kicker_y
        title_y = GEOMETRY.dark_title_y if is_dark else GEOMETRY.header_title_y

        # 1. Icon Badge
        badge = IconBadge.render(
            slide,
            x=GEOMETRY.header_badge_x,
            y=badge_y,
            diameter=GEOMETRY.header_badge_d,
            icon_path=icon_path,
            semantic_hint=semantic_hint,
            slide_idx=slide_idx
        )

        # 2. Kicker Text
        clean_kicker = format_kicker_text(kicker)
        tb_kicker = add_text_box(slide, GEOMETRY.header_kicker_x, kicker_y, 11.20, 0.28)
        pk = tb_kicker.text_frame.paragraphs[0]
        pk.text = clean_kicker
        pk.font.name = FONTS.body
        pk.font.size = Pt(12.0)
        pk.font.bold = True
        pk.font.color.rgb = COLORS.rgb_kicker

        # 3. Slide Title
        clean_title, _ = TitleOptimizer.optimize(title, max_words=8)
        tb_title = add_text_box(slide, GEOMETRY.header_title_x, title_y, GEOMETRY.header_title_w, 0.68)
        pt = tb_title.text_frame.paragraphs[0]
        pt.text = clean_title
        pt.font.name = FONTS.display
        pt.font.size = Pt(27.0)
        pt.font.bold = True
        pt.font.color.rgb = COLORS.rgb_white if is_dark else COLORS.rgb_heading

        return badge, tb_kicker, tb_title


class SlideFooter:
    """Renders the standard footer: presentation/source line on left, page number on right."""

    @classmethod
    def render(
        cls,
        slide: Any,
        slide_num: int,
        source_text: str = "The Rise of Empires  •  NCERT | Exploring Society: India and Beyond",
        is_dark: bool = False
    ):
        y = GEOMETRY.footer_top
        # Left source line
        tb_left = add_text_box(slide, GEOMETRY.footer_left, y, 10.50, 0.28)
        pl = tb_left.text_frame.paragraphs[0]
        pl.text = source_text
        pl.font.name = FONTS.body
        pl.font.size = Pt(9.0)
        pl.font.color.rgb = COLORS.rgb_text_on_dark if is_dark else COLORS.rgb_muted

        # Right page number
        tb_right = add_text_box(slide, GEOMETRY.footer_slide_num_x, y, 0.60, 0.28)
        pr = tb_right.text_frame.paragraphs[0]
        pr.text = str(slide_num)
        pr.alignment = PP_ALIGN.RIGHT
        pr.font.name = FONTS.body
        pr.font.size = Pt(9.0)
        pr.font.color.rgb = COLORS.rgb_text_on_dark if is_dark else COLORS.rgb_muted


class ParchmentCard:
    """Renders a light parchment information card (#FAF6EF or #F4EEE4) with subtle border."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        title: Optional[str] = None,
        body_bullets: Optional[List[str]] = None,
        fill_color: RGBColor = COLORS.rgb_parchment,
        border_color: RGBColor = COLORS.rgb_border_warm,
        shadow: bool = True,
        title_size_pt: float = 16.0,
        body_size_pt: float = 14.0
    ) -> Any:
        card = create_solid_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            y,
            w,
            h,
            fill_color=fill_color,
            border_color=border_color,
            shadow=shadow
        )

        pad_x = 0.28
        pad_y = 0.24
        inner_w = w - (2 * pad_x)
        inner_h = h - (2 * pad_y)

        if title or body_bullets:
            tb = add_text_box(slide, x + pad_x, y + pad_y, inner_w, inner_h)
            tf = tb.text_frame
            tf.word_wrap = True

            first = True
            if title:
                p0 = tf.paragraphs[0]
                p0.text = title
                p0.font.name = FONTS.display
                p0.font.size = Pt(title_size_pt)
                p0.font.bold = True
                p0.font.color.rgb = COLORS.rgb_heading
                p0.space_after = Pt(8)
                first = False

            if body_bullets:
                for b_text in body_bullets:
                    clean_b = b_text.strip()
                    if not clean_b:
                        continue
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    p.text = clean_b
                    p.font.name = FONTS.body
                    p.font.size = Pt(body_size_pt)
                    p.font.color.rgb = COLORS.rgb_body
                    p.space_after = Pt(6)
                    p.line_spacing = 1.15
                    if len(body_bullets) > 1:
                        p.level = 0

        return card


class DarkInsightCard:
    """Renders a deep teal (#123B47) insight card with light warm cream text."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        title: Optional[str] = None,
        body_text: Optional[str] = None,
        title_size_pt: float = 16.0,
        body_size_pt: float = 13.5
    ) -> Any:
        card = create_solid_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            y,
            w,
            h,
            fill_color=COLORS.rgb_canvas_dark,
            border_color=None,
            shadow=True
        )

        pad_x = 0.28
        pad_y = 0.24
        tb = add_text_box(slide, x + pad_x, y + pad_y, w - (2 * pad_x), h - (2 * pad_y))
        tf = tb.text_frame
        tf.word_wrap = True

        first = True
        if title:
            p0 = tf.paragraphs[0]
            p0.text = title
            p0.font.name = FONTS.display
            p0.font.size = Pt(title_size_pt)
            p0.font.bold = True
            p0.font.color.rgb = COLORS.rgb_white
            p0.space_after = Pt(6)
            first = False

        if body_text:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.text = body_text
            p.font.name = FONTS.body
            p.font.size = Pt(body_size_pt)
            p.font.color.rgb = COLORS.rgb_text_on_dark
            p.line_spacing = 1.15

        return card


class GlossaryStrip:
    """Renders a single-frame glossary callout in #123B47 with burnt orange term."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        term: str,
        definition: str
    ) -> Any:
        panel = create_solid_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            y,
            w,
            h,
            fill_color=COLORS.rgb_canvas_dark,
            border_color=None,
            shadow=False
        )
        tf = panel.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.24)
        tf.margin_right = Inches(0.24)
        tf.margin_top = Inches(0.10)
        tf.margin_bottom = Inches(0.10)

        p = tf.paragraphs[0]
        p.line_spacing = 1.15

        # Term run (Burnt Orange, Bold Cambria)
        r1 = p.add_run()
        r1.text = f"{term}   "
        r1.font.name = FONTS.display
        r1.font.size = Pt(13.0)
        r1.font.bold = True
        r1.font.color.rgb = COLORS.rgb_accent

        # Definition run (Light Warm Cream Calibri)
        r2 = p.add_run()
        r2.text = definition
        r2.font.name = FONTS.body
        r2.font.size = Pt(12.0)
        r2.font.color.rgb = COLORS.rgb_text_on_dark

        return panel


class HistoricalImageFrame:
    """Renders an image inside a warm cream frame with caption underneath."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        image_path: str,
        caption: Optional[str] = None,
        caption_h: float = 0.45
    ) -> Tuple[Any, Optional[Any]]:
        frame_h = h - caption_h if caption else h
        frame = create_solid_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            y,
            w,
            frame_h,
            fill_color=COLORS.rgb_cream,
            border_color=COLORS.rgb_border_warm,
            shadow=True
        )

        # Image padding
        pad = 0.08
        img_x = x + pad
        img_y = y + pad
        img_w = w - (2 * pad)
        img_h = frame_h - (2 * pad)

        if os.path.exists(image_path):
            try:
                slide.shapes.add_picture(image_path, Inches(img_x), Inches(img_y), Inches(img_w), Inches(img_h))
            except Exception as e:
                logger.warning("Could not add picture %s: %s", image_path, e)

        tb_caption = None
        if caption:
            tb_caption = add_text_box(slide, x, y + frame_h + 0.06, w, caption_h - 0.06)
            pc = tb_caption.text_frame.paragraphs[0]
            pc.text = caption
            pc.font.name = FONTS.body
            pc.font.size = Pt(10.5)
            pc.font.italic = True
            pc.font.color.rgb = COLORS.rgb_muted

        return frame, tb_caption


class HistoricalQuote:
    """Renders historical quote with large burnt-orange quotation mark and Cambria italic text."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        quote_text: str,
        attribution: Optional[str] = None,
        is_dark: bool = True
    ) -> Any:
        tb = add_text_box(slide, x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True

        # Quotation mark
        p0 = tf.paragraphs[0]
        p0.text = "“"
        p0.font.name = FONTS.display
        p0.font.size = Pt(48.0)
        p0.font.bold = True
        p0.font.color.rgb = COLORS.rgb_accent
        p0.space_after = Pt(2)

        # Quote body
        p1 = tf.add_paragraph()
        p1.text = quote_text
        p1.font.name = FONTS.display
        p1.font.size = Pt(20.0)
        p1.font.italic = True
        p1.font.color.rgb = COLORS.rgb_white if is_dark else COLORS.rgb_heading
        p1.space_after = Pt(12)
        p1.line_spacing = 1.2

        # Attribution
        if attribution:
            p2 = tf.add_paragraph()
            p2.text = f"— {attribution}"
            p2.font.name = FONTS.body
            p2.font.size = Pt(12.0)
            p2.font.color.rgb = COLORS.rgb_kicker if is_dark else COLORS.rgb_muted

        return tb


class HubAndSpokeDiagram:
    """Renders central emperor card connected to 6 peripheral numbered feature cards."""

    @classmethod
    def render(
        cls,
        slide: Any,
        center_title: str,
        center_body: str,
        features: List[Tuple[str, str]],  # (title, body) up to 6
        center_icon_path: Optional[str] = None
    ):
        # 1. Left 3 Cards
        card_w = 3.55
        card_h = 1.28
        left_x = 0.55
        right_x = 9.23
        row_ys = [1.72, 3.15, 4.58]

        for i in range(min(3, len(features))):
            f_title, f_desc = features[i]
            y = row_ys[i]
            # Card
            ParchmentCard.render(slide, left_x, y, card_w, card_h, fill_color=COLORS.rgb_parchment)
            # Number circle
            NumberBadge.render(slide, left_x + 0.20, y + 0.31, str(i + 1), diameter=0.66)
            # Text box
            tb = add_text_box(slide, left_x + 1.00, y + 0.14, card_w - 1.15, card_h - 0.28)
            tf = tb.text_frame
            tf.word_wrap = True
            p0 = tf.paragraphs[0]
            p0.text = f_title
            p0.font.name = FONTS.display
            p0.font.size = Pt(14.0)
            p0.font.bold = True
            p0.font.color.rgb = COLORS.rgb_heading
            p0.space_after = Pt(4)
            p1 = tf.add_paragraph()
            p1.text = f_desc
            p1.font.name = FONTS.body
            p1.font.size = Pt(11.5)
            p1.font.color.rgb = COLORS.rgb_body
            p1.line_spacing = 1.1

        # 2. Right 3 Cards
        for i in range(3, min(6, len(features))):
            f_title, f_desc = features[i]
            idx = i - 3
            y = row_ys[idx]
            # Card
            ParchmentCard.render(slide, right_x, y, card_w, card_h, fill_color=COLORS.rgb_parchment)
            # Number circle
            NumberBadge.render(slide, right_x + 0.20, y + 0.31, str(i + 1), diameter=0.66)
            # Text box
            tb = add_text_box(slide, right_x + 1.00, y + 0.14, card_w - 1.15, card_h - 0.28)
            tf = tb.text_frame
            tf.word_wrap = True
            p0 = tf.paragraphs[0]
            p0.text = f_title
            p0.font.name = FONTS.display
            p0.font.size = Pt(14.0)
            p0.font.bold = True
            p0.font.color.rgb = COLORS.rgb_heading
            p0.space_after = Pt(4)
            p1 = tf.add_paragraph()
            p1.text = f_desc
            p1.font.name = FONTS.body
            p1.font.size = Pt(11.5)
            p1.font.color.rgb = COLORS.rgb_body
            p1.line_spacing = 1.1

        # 3. Central Dark Emperor Shape
        cx = 5.32
        cy = 2.90
        cw = 2.70
        ch = 2.30
        c_card = create_solid_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            cx,
            cy,
            cw,
            ch,
            fill_color=COLORS.rgb_canvas_dark,
            border_color=None,
            shadow=True
        )

        # Center icon
        if center_icon_path and os.path.exists(center_icon_path):
            slide.shapes.add_picture(center_icon_path, Inches(cx + 0.95), Inches(cy + 0.25), Inches(0.80), Inches(0.80))
            text_y = cy + 1.05
            text_h = ch - 1.15
        else:
            text_y = cy + 0.30
            text_h = ch - 0.50

        tb_c = add_text_box(slide, cx + 0.15, text_y, cw - 0.30, text_h)
        tfc = tb_c.text_frame
        tfc.word_wrap = True
        p0 = tfc.paragraphs[0]
        p0.alignment = PP_ALIGN.CENTER
        p0.text = center_title
        p0.font.name = FONTS.display
        p0.font.size = Pt(16.0)
        p0.font.bold = True
        p0.font.color.rgb = COLORS.rgb_white
        p0.space_after = Pt(6)

        p1 = tfc.add_paragraph()
        p1.alignment = PP_ALIGN.CENTER
        p1.text = center_body
        p1.font.name = FONTS.body
        p1.font.size = Pt(12.0)
        p1.font.color.rgb = COLORS.rgb_text_on_dark
        p1.line_spacing = 1.15


class HistoricalTimeline:
    """Renders chronological historical timeline with antique gold track and era cards."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        w: float,
        eras: List[Tuple[str, str, str]],  # (era_name, dates, description)
    ):
        n = len(eras)
        if n == 0:
            return

        # Antique gold horizontal baseline track
        line_y = y + 2.40
        track = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(line_y), Inches(w), Inches(0.06))
        track.fill.solid()
        track.fill.fore_color.rgb = COLORS.rgb_kicker
        track.line.fill.background()

        card_gap = 0.35
        card_w = (w - (n - 1) * card_gap) / n

        for i, (era_name, dates, desc) in enumerate(eras):
            cx = x + i * (card_w + card_gap)
            node_center_x = cx + card_w / 2.0

            # Era Card above baseline
            ParchmentCard.render(
                slide,
                cx,
                y,
                card_w,
                2.15,
                title=era_name,
                body_bullets=[dates, desc],
                fill_color=COLORS.rgb_parchment,
                title_size_pt=15.0,
                body_size_pt=12.5
            )

            # Burnt orange milestone node circle on track
            node_d = 0.40
            node = create_solid_shape(
                slide,
                MSO_SHAPE.OVAL,
                node_center_x - (node_d / 2.0),
                line_y - (node_d / 2.0) + 0.03,
                node_d,
                node_d,
                fill_color=COLORS.rgb_accent,
                border_color=COLORS.rgb_white,
                border_width_pt=1.5
            )


class TakeawayRow:
    """Renders structured legacy takeaways with burnt orange circular markers."""

    @classmethod
    def render(
        cls,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        items: List[str],
        is_dark: bool = True
    ):
        row_h = h / max(1, len(items))
        for i, text in enumerate(items):
            iy = y + i * row_h
            # Orange check / dot marker
            marker_d = 0.32
            NumberBadge.render(slide, x, iy + 0.08, str(i + 1), diameter=marker_d, font_size_pt=9.0)

            # Text box
            tb = add_text_box(slide, x + marker_d + 0.20, iy, w - marker_d - 0.20, row_h - 0.08)
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = text
            p.font.name = FONTS.body
            p.font.size = Pt(14.0)
            p.font.color.rgb = COLORS.rgb_text_on_dark if is_dark else COLORS.rgb_body
            p.line_spacing = 1.15
