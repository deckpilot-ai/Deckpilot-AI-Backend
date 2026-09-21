"""Layout Archetype Library for Maratha Heritage Editorial Theme (M01 - M24).

Fully implements all 24 reference layout builders with strict grid alignment,
zero text clipping, responsive typography, and authentic heritage visual hierarchy.
"""

import logging
import os
from typing import Any, Dict, List, Optional, Tuple

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.presentation.themes.maratha_heritage_editorial.tokens import COLORS, GEOMETRY, FONTS
from app.presentation.themes.maratha_heritage_editorial.components import (
    create_solid_shape, add_text_box, MarathaEditorialHeader, SlideFooter,
    IconBadge, NumberBadge, CheckBadge, ParchmentCard, MaroonInsightCard,
    MaroonInsightStrip, DidYouKnowCard, HeritageGlossaryCard, HistoricalImageFrame,
    HistoricalMetricCard, MarathaHistoricalQuote, TimelinePhaseCard,
    GovernanceHubDiagram, ThreePillarCard, DualBiographyCard, MarathaLegacySummary
)
from app.presentation.themes.maratha_heritage_editorial.icons import IconResolver

logger = logging.getLogger(__name__)

ASSETS_DIR = os.path.join(os.path.dirname(__file__), "assets")
IMAGES_DIR = os.path.join(ASSETS_DIR, "images")
ICONS_DIR = os.path.join(ASSETS_DIR, "icons")


def get_image_asset(filename: str) -> str:
    path = os.path.join(IMAGES_DIR, filename)
    if os.path.exists(path):
        return path
    ipath = os.path.join(ICONS_DIR, filename)
    if os.path.exists(ipath):
        return ipath
    return path


def set_slide_background(slide: Any, color: RGBColor):
    """Fills slide background solidly."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


# ====================================================================
# M01 — MAROON COVER + HERO IMAGE
# ====================================================================
def render_m01_maroon_cover(
    slide: Any,
    kicker: str,
    title: str,
    subtitle: Optional[str] = None,
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 1
):
    set_slide_background(slide, COLORS.rgb_maroon)

    # Subtle decorative circular geometry
    create_solid_shape(slide, MSO_SHAPE.OVAL, -1.2, -1.2, 3.8, 3.8, fill_color=COLORS.rgb_maroon_alt)
    create_solid_shape(slide, MSO_SHAPE.OVAL, 10.8, 4.8, 4.2, 4.2, fill_color=COLORS.rgb_maroon_alt)

    # Gold Kicker
    tb_k = add_text_box(slide, 0.55, 1.20, 7.00, 0.35)
    pk = tb_k.text_frame.paragraphs[0]
    pk.text = kicker.upper()
    pk.font.name = FONTS.body
    pk.font.size = Pt(12.0)
    pk.font.bold = True
    pk.font.color.rgb = COLORS.rgb_gold

    # Title (Large White Cambria, font size and vertical layout dynamically responsive)
    if len(title) > 55:
        title_font_size = 38.0
        line_offset = 4.05
    elif len(title) > 35:
        title_font_size = 42.0
        line_offset = 3.85
    else:
        title_font_size = 50.0
        line_offset = 3.65

    tb_t = add_text_box(slide, 0.55, 1.65, 7.00, line_offset - 1.70)
    tf_t = tb_t.text_frame
    pt = tf_t.paragraphs[0]
    pt.text = title
    pt.font.name = FONTS.display
    pt.font.size = Pt(title_font_size)
    pt.font.bold = True
    pt.font.color.rgb = COLORS.rgb_white
    pt.line_spacing = 1.05

    # Orange horizontal divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(line_offset), Inches(1.50), Inches(0.055))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS.rgb_orange
    line.line.fill.background()

    # Subtitle / Chapter metadata
    if subtitle:
        tb_s = add_text_box(slide, 0.55, line_offset + 0.20, 6.80, 1.20)
        ps = tb_s.text_frame.paragraphs[0]
        ps.text = subtitle
        ps.font.name = FONTS.body
        ps.font.size = Pt(14.0)
        ps.font.italic = True
        ps.font.color.rgb = COLORS.rgb_text_on_dark

    # Hero Image Right
    img = image_path or get_image_asset("image-1-1.png")
    HistoricalImageFrame.render(
        slide,
        x=7.82,
        y=1.568,
        w=5.01,
        h=4.80,
        image_path=img,
        caption=caption or "Coronation of Chhatrapati Shivaji",
        caption_h=0.45
    )


# ====================================================================
# M02 — THREE BIG QUESTIONS
# ====================================================================
def render_m02_three_questions(
    slide: Any,
    kicker: str,
    title: str,
    questions: List[str],
    slide_num: int = 2
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=2)
    SlideFooter.render(slide, slide_num=slide_num)

    n = max(1, min(4, len(questions)))
    card_gap = 0.35
    total_w = 12.233
    card_w = (total_w - (n - 1) * card_gap) / n
    card_h = 4.80
    top_y = 1.75

    for i, q in enumerate(questions[:n]):
        cx = 0.55 + i * (card_w + card_gap)
        # Parchment card
        ParchmentCard.render(slide, cx, top_y, card_w, card_h, fill_color=COLORS.rgb_parchment)

        # Large orange sequence circle
        num_d = 0.72
        NumberBadge.render(slide, cx + (card_w - num_d) / 2.0, top_y + 0.45, str(i + 1), diameter=num_d, font_size_pt=16.0)

        # Question text centered
        tb_q = add_text_box(slide, cx + 0.25, top_y + 1.45, card_w - 0.50, card_h - 1.80)
        tf_q = tb_q.text_frame
        tf_q.word_wrap = True
        tf_q.vertical_anchor = MSO_ANCHOR.MIDDLE
        pq = tf_q.paragraphs[0]
        pq.alignment = PP_ALIGN.CENTER
        pq.text = q
        pq.font.name = FONTS.display
        pq.font.size = Pt(17.5)
        pq.font.bold = True
        pq.font.color.rgb = COLORS.rgb_maroon
        pq.line_spacing = 1.25


# ====================================================================
# M03 — ORIGIN STORY + TWO VISUALS + GLOSSARY
# ====================================================================
def render_m03_origin_glossary(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    glossary_term: str,
    glossary_def: str,
    image1_path: Optional[str] = None,
    image2_path: Optional[str] = None,
    caption1: Optional[str] = None,
    caption2: Optional[str] = None,
    slide_num: int = 3
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=3)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.75,
        w=8.20,
        h=4.80,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Two Historical Visuals Upper Right
    img1 = image1_path or get_image_asset("image-3-2.png")
    HistoricalImageFrame.render(
        slide,
        x=9.05,
        y=1.75,
        w=1.85,
        h=3.20,
        image_path=img1,
        caption=caption1 or "Shahaji Bhonsle",
        caption_h=0.40
    )

    img2 = image2_path or get_image_asset("image-3-3.png")
    HistoricalImageFrame.render(
        slide,
        x=11.15,
        y=2.00,
        w=1.75,
        h=2.95,
        image_path=img2,
        caption=caption2 or "Jijabai",
        caption_h=0.40
    )

    # Glossary Card Lower Right
    HeritageGlossaryCard.render(
        slide,
        x=9.05,
        y=5.25,
        w=3.85,
        h=1.30,
        term=glossary_term,
        definition=glossary_def,
        icon_path=get_image_asset("image-3-1.png"),
        is_maroon=True
    )


# ====================================================================
# M04 — THREE-PHASE CHRONOLOGY
# ====================================================================
def render_m04_three_phase_timeline(
    slide: Any,
    kicker: str,
    title: str,
    phases: List[Dict[str, Any]],  # [{title, date_range, events: [(date, desc)]}]
    slide_num: int = 4
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=4)
    SlideFooter.render(slide, slide_num=slide_num)

    col_w = 3.84
    col_gap = 0.35
    col_h = 4.80

    for i, p in enumerate(phases[:3]):
        cx = 0.55 + i * (col_w + col_gap)
        TimelinePhaseCard.render(
            slide,
            x=cx,
            y=1.75,
            w=col_w,
            h=col_h,
            phase_title=p.get("title", f"PHASE {i+1}"),
            date_range=p.get("date_range", ""),
            events=p.get("events", [])
        )


# ====================================================================
# M05 — BIOGRAPHY + PORTRAIT + KEY TERM
# ====================================================================
def render_m05_biography_glossary(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    term: str,
    definition: str,
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 5
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=5)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=8.50,
        h=3.50,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Wide Key-Term Card below narrative
    HeritageGlossaryCard.render(
        slide,
        x=0.55,
        y=5.45,
        w=8.50,
        h=1.10,
        term=term,
        definition=definition,
        icon_path=get_image_asset("image-5-3.png"),
        is_maroon=True
    )

    # Hero Portrait Right
    img = image_path or get_image_asset("image-5-2.png")
    HistoricalImageFrame.render(
        slide,
        x=9.40,
        y=1.72,
        w=3.40,
        h=4.80,
        image_path=img,
        caption=caption or "Chhatrapati Shivaji Maharaj",
        caption_h=0.45
    )


# ====================================================================
# M06 — EXPLANATION + IMAGE + TWO METRICS
# ====================================================================
def render_m06_narrative_metrics_image(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    metrics: List[Tuple[str, str, Optional[str]]],  # [(val, label, subtitle)]
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 6
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=6)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative card left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.00,
        h=3.20,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Two Metric cards bottom-left
    m_w = 3.35
    for i, m in enumerate(metrics[:2]):
        mx = 0.55 + i * (m_w + 0.30)
        HistoricalMetricCard.render(
            slide,
            x=mx,
            y=5.10,
            w=m_w,
            h=1.45,
            metric_value=m[0],
            label=m[1],
            subtitle=m[2] if len(m) > 2 else None
        )

    # Large Historical Photo Right
    img = image_path or get_image_asset("image-6-2.png")
    HistoricalImageFrame.render(
        slide,
        x=7.90,
        y=1.72,
        w=4.90,
        h=4.80,
        image_path=img,
        caption=caption or "Sindhudurg sea fort, fortified naval base",
        caption_h=0.45
    )


# ====================================================================
# M07 — EXPLANATION + ARTEFACT + DID-YOU-KNOW
# ====================================================================
def render_m07_artefact_didyouknow(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    did_you_know_text: str,
    artefact_image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 7
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=7)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.40,
        h=4.80,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Artefact Image Top-Right
    img = artefact_image_path or get_image_asset("image-7-2.png")
    HistoricalImageFrame.render(
        slide,
        x=8.30,
        y=1.72,
        w=4.50,
        h=2.40,
        image_path=img,
        caption=caption or "Waghnakh (tiger claws) concealed weapon",
        caption_h=0.40
    )

    # Did-You-Know Card Bottom-Right
    DidYouKnowCard.render(
        slide,
        x=8.30,
        y=4.35,
        w=4.50,
        h=2.15,
        body_text=did_you_know_text
    )


# ====================================================================
# M08 — TWO EVENT CARDS + INSIGHT STRIP
# ====================================================================
def render_m08_dual_campaign_cards(
    slide: Any,
    kicker: str,
    title: str,
    card1_title: str,
    card1_body: str,
    card2_title: str,
    card2_body: str,
    insight_text: str,
    slide_num: int = 8
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=8)
    SlideFooter.render(slide, slide_num=slide_num)

    card_w = 5.95
    card_h = 3.65
    top_y = 1.72

    # Card 1
    c1 = ParchmentCard.render(slide, 0.55, top_y, card_w, card_h, fill_color=COLORS.rgb_parchment)
    IconBadge.render(slide, 0.85, top_y + 0.30, diameter=0.60, icon_path=get_image_asset("image-8-2.png"))
    tb1 = add_text_box(slide, 1.60, top_y + 0.25, card_w - 1.20, 0.45)
    p1 = tb1.text_frame.paragraphs[0]
    p1.text = card1_title
    p1.font.name = FONTS.display
    p1.font.size = Pt(17.0)
    p1.font.bold = True
    p1.font.color.rgb = COLORS.rgb_maroon

    tb1_b = add_text_box(slide, 0.85, top_y + 1.05, card_w - 0.60, card_h - 1.25)
    tf1_b = tb1_b.text_frame
    tf1_b.word_wrap = True
    pb1 = tf1_b.paragraphs[0]
    pb1.text = card1_body
    pb1.font.name = FONTS.body
    pb1.font.size = Pt(13.0)
    pb1.font.color.rgb = COLORS.rgb_body
    pb1.line_spacing = 1.15

    # Card 2
    c2 = ParchmentCard.render(slide, 6.85, top_y, card_w, card_h, fill_color=COLORS.rgb_parchment)
    IconBadge.render(slide, 7.15, top_y + 0.30, diameter=0.60, icon_path=get_image_asset("image-8-3.png"))
    tb2 = add_text_box(slide, 7.90, top_y + 0.25, card_w - 1.20, 0.45)
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = card2_title
    p2.font.name = FONTS.display
    p2.font.size = Pt(17.0)
    p2.font.bold = True
    p2.font.color.rgb = COLORS.rgb_maroon

    tb2_b = add_text_box(slide, 7.15, top_y + 1.05, card_w - 0.60, card_h - 1.25)
    tf2_b = tb2_b.text_frame
    tf2_b.word_wrap = True
    pb2 = tf2_b.paragraphs[0]
    pb2.text = card2_body
    pb2.font.name = FONTS.body
    pb2.font.size = Pt(13.0)
    pb2.font.color.rgb = COLORS.rgb_body
    pb2.line_spacing = 1.15

    # Full-Width Maroon Insight Strip
    MaroonInsightStrip.render(
        slide,
        x=0.55,
        y=5.60,
        w=12.233,
        h=0.95,
        text=insight_text,
        highlight_prefix="IMPACT:",
        icon_path=get_image_asset("image-8-1.png")
    )


# ====================================================================
# M09 — WIDE STORY PANEL + CENTERED HISTORICAL IMAGE
# ====================================================================
def render_m09_story_centered_image(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 9
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=9)
    SlideFooter.render(slide, slide_num=slide_num)

    # Wide Narrative Card
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=12.233,
        h=2.65,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Centered Horizontal Historical Image Beneath
    img = image_path or get_image_asset("image-9-2.png")
    HistoricalImageFrame.render(
        slide,
        x=3.80,
        y=4.55,
        w=5.733,
        h=2.00,
        image_path=img,
        caption=caption or "Shivaji escaping Mughal captivity in Agra concealed in sweet baskets",
        caption_h=0.38
    )


# ====================================================================
# M10 — NARRATIVE + MAP + DATE CALLOUT
# ====================================================================
def render_m10_narrative_map(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    date_callout: str,
    callout_label: str,
    map_image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 10
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=10)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=8.10,
        h=3.50,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Date Callout Strip
    HistoricalMetricCard.render(
        slide,
        x=0.55,
        y=5.40,
        w=8.10,
        h=1.15,
        metric_value=date_callout,
        label=callout_label
    )

    # Tall Map Right
    img = map_image_path or get_image_asset("image-10-2.png")
    HistoricalImageFrame.render(
        slide,
        x=8.95,
        y=1.72,
        w=3.833,
        h=4.80,
        image_path=img,
        caption=caption or "Maratha expansion and southern campaigns",
        caption_h=0.45
    )


# ====================================================================
# M11 — NARRATIVE + DARK ETHICAL / INSIGHT PANEL
# ====================================================================
def render_m11_narrative_insight_panel(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    panel_title: str,
    panel_body: str,
    slide_num: int = 11
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=11)
    SlideFooter.render(slide, slide_num=slide_num)

    # Large Parchment Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.40,
        h=4.80,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Tall Dark-Maroon Contextual Panel Right
    MaroonInsightCard.render(
        slide,
        x=8.25,
        y=1.72,
        w=4.533,
        h=4.80,
        title=panel_title,
        body_text=panel_body,
        icon_path=get_image_asset("image-11-2.png"),
        title_size_pt=17.0,
        body_size_pt=13.0
    )


# ====================================================================
# M12 — NARRATIVE + HISTORICAL PAINTING
# ====================================================================
def render_m12_narrative_painting(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 12
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=12)
    SlideFooter.render(slide, slide_num=slide_num)

    # Large narrative panel left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=8.40,
        h=4.80,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Tall Historical Artwork Right
    img = image_path or get_image_asset("image-12-2.png")
    HistoricalImageFrame.render(
        slide,
        x=9.25,
        y=1.72,
        w=3.533,
        h=4.80,
        image_path=img,
        caption=caption or "Chhatrapati Sambhaji Maharaj",
        caption_h=0.45
    )


# ====================================================================
# M13 — NARRATIVE + IMAGE + FULL-WIDTH STRIP
# ====================================================================
def render_m13_narrative_strip(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    strip_title: str,
    strip_body: str,
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 13
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=13)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.50,
        h=3.50,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Image Right
    img = image_path or get_image_asset("image-13-2.png")
    HistoricalImageFrame.render(
        slide,
        x=8.35,
        y=1.72,
        w=4.433,
        h=3.50,
        image_path=img,
        caption=caption or "Peshwa Baji Rao I",
        caption_h=0.40
    )

    # Bottom Full-Width Maroon Strip
    MaroonInsightStrip.render(
        slide,
        x=0.55,
        y=5.45,
        w=12.233,
        h=1.10,
        text=strip_body,
        highlight_prefix=strip_title,
        icon_path=get_image_asset("image-13-3.png")
    )


# ====================================================================
# M14 — THREE WAR / PERIOD CARDS + QUOTE
# ====================================================================
def render_m14_three_periods_quote(
    slide: Any,
    kicker: str,
    title: str,
    periods: List[Dict[str, str]],  # [{date, title, desc}]
    quote_text: str,
    attribution: Optional[str] = None,
    slide_num: int = 14
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=14)
    SlideFooter.render(slide, slide_num=slide_num)

    card_w = 3.84
    card_gap = 0.35
    card_h = 3.50
    top_y = 1.72

    for i, p in enumerate(periods[:3]):
        cx = 0.55 + i * (card_w + card_gap)
        ParchmentCard.render(
            slide,
            cx,
            top_y,
            card_w,
            card_h,
            title=p.get("title", ""),
            body_bullets=[p.get("date", ""), p.get("desc", "")],
            fill_color=COLORS.rgb_parchment,
            title_size_pt=15.0,
            body_size_pt=12.5
        )

    # Full-Width Dark Quote Strip
    panel = create_solid_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 5.45, 12.233, 1.10, fill_color=COLORS.rgb_maroon)
    tb = add_text_box(slide, 1.10, 5.50, 11.50, 1.00)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pq = tf.paragraphs[0]
    pq.text = f"“{quote_text}”"
    if attribution:
        pq.text += f"  — {attribution}"
    pq.font.name = FONTS.display
    pq.font.size = Pt(13.0)
    pq.font.italic = True
    pq.font.color.rgb = COLORS.rgb_white


# ====================================================================
# M15 — ADMINISTRATION + COIN + BOTTOM INSIGHT
# ====================================================================
def render_m15_narrative_coin_insight(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    insight_text: str,
    coin_image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 15
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=15)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.40,
        h=3.50,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Coin Image Right
    img = coin_image_path or get_image_asset("image-15-2.png")
    HistoricalImageFrame.render(
        slide,
        x=8.25,
        y=1.72,
        w=4.533,
        h=3.50,
        image_path=img,
        caption=caption or "Hon (gold coin) minted by Chhatrapati Shivaji",
        caption_h=0.40
    )

    # Bottom Cream Insight Strip
    ParchmentCard.render(
        slide,
        x=0.55,
        y=5.45,
        w=12.233,
        h=1.10,
        body_bullets=[insight_text],
        fill_color=COLORS.rgb_cream,
        body_size_pt=12.5
    )


# ====================================================================
# M16 — GOVERNANCE HUB / COUNCIL OF EIGHT
# ====================================================================
def render_m16_governance_hub(
    slide: Any,
    kicker: str,
    title: str,
    center_title: str,
    center_subtitle: str,
    top_ministers: List[Tuple[str, str, str]],
    bottom_ministers: List[Tuple[str, str, str]],
    slide_num: int = 16
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=16)
    SlideFooter.render(slide, slide_num=slide_num)

    GovernanceHubDiagram.render(
        slide,
        x=0.55,
        y=1.65,
        w=12.233,
        h=4.90,
        center_title=center_title,
        center_subtitle=center_subtitle,
        top_cards=top_ministers,
        bottom_cards=bottom_ministers
    )


# ====================================================================
# M17 — REVENUE + COIN + KPI CARDS
# ====================================================================
def render_m17_revenue_kpi_coin(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    kpis: List[Tuple[str, str, Optional[str]]],
    coin_image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 17
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=17)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.30,
        h=3.20,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Two KPI Cards Bottom Left
    kpi_w = 3.50
    for i, k in enumerate(kpis[:2]):
        kx = 0.55 + i * (kpi_w + 0.30)
        HistoricalMetricCard.render(
            slide,
            x=kx,
            y=5.10,
            w=kpi_w,
            h=1.45,
            metric_value=k[0],
            label=k[1],
            subtitle=k[2] if len(k) > 2 else None
        )

    # Coin Image Right
    img = coin_image_path or get_image_asset("image-17-2.png")
    HistoricalImageFrame.render(
        slide,
        x=8.15,
        y=1.72,
        w=4.633,
        h=4.80,
        image_path=img,
        caption=caption or "Maratha rupee coin with Devnagari inscription",
        caption_h=0.45
    )


# ====================================================================
# M18 — MILITARY SYSTEM + WEAPONS IMAGE
# ====================================================================
def render_m18_military_system(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 18
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=18)
    SlideFooter.render(slide, slide_num=slide_num)

    # Large Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.40,
        h=4.80,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Weapon / Equipment Image Right
    img = image_path or get_image_asset("image-18-2.png")
    HistoricalImageFrame.render(
        slide,
        x=8.25,
        y=1.72,
        w=4.533,
        h=4.80,
        image_path=img,
        caption=caption or "Maratha Firangi straight sword and cavalry shield",
        caption_h=0.45
    )


# ====================================================================
# M19 — DARK QUOTE + THREE SUPPORTING POINTS
# ====================================================================
def render_m19_dark_quote(
    slide: Any,
    kicker: str,
    title: str,
    quote_text: str,
    attribution: str,
    supporting_points: List[str],
    slide_num: int = 19
):
    set_slide_background(slide, COLORS.rgb_maroon)
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=19, is_dark=True)
    SlideFooter.render(slide, slide_num=slide_num, is_dark=True)

    # Large Quote on Left
    MarathaHistoricalQuote.render(
        slide,
        x=0.55,
        y=1.85,
        w=5.80,
        h=4.60,
        quote_text=quote_text,
        attribution=attribution,
        is_dark=True
    )

    # Three Vertically Stacked Principles Right
    col_x = 6.85
    col_w = 5.933
    row_h = 1.35
    for i, pt_text in enumerate(supporting_points[:3]):
        iy = 1.85 + i * (row_h + 0.25)
        # Check / Marker Icon
        CheckBadge.render(slide, col_x, iy + 0.10, size=0.34)
        tb = add_text_box(slide, col_x + 0.50, iy, col_w - 0.50, row_h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = pt_text
        p.font.name = FONTS.body
        p.font.size = Pt(13.5)
        p.font.color.rgb = COLORS.rgb_text_on_dark
        p.line_spacing = 1.15


# ====================================================================
# M20 — MARITIME + IMAGE + DARK FACT STRIP
# ====================================================================
def render_m20_maritime_fact(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    fact_text: str,
    image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 20
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=20)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=6.90,
        h=4.80,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Historical Maritime Image Right
    img = image_path or get_image_asset("image-20-2.png")
    HistoricalImageFrame.render(
        slide,
        x=7.75,
        y=1.72,
        w=5.033,
        h=3.40,
        image_path=img,
        caption=caption or "Maratha warships engaging European vessels off Konkan coast",
        caption_h=0.40
    )

    # Dark Contextual Fact Strip below image
    MaroonInsightStrip.render(
        slide,
        x=7.75,
        y=5.35,
        w=5.033,
        h=1.15,
        text=fact_text,
        highlight_prefix="KEY OUTCOME:"
    )


# ====================================================================
# M21 — THREE PILLARS + SUPPORTING SEAL
# ====================================================================
def render_m21_three_pillars_seal(
    slide: Any,
    kicker: str,
    title: str,
    pillars: List[Tuple[str, str, str]],  # [(title, body, icon_path)]
    seal_image_path: Optional[str] = None,
    caption: Optional[str] = None,
    slide_num: int = 21
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=21)
    SlideFooter.render(slide, slide_num=slide_num)

    col_w = 2.767
    col_gap = 0.35
    col_h = 4.80

    for i, p in enumerate(pillars[:3]):
        cx = 0.55 + i * (col_w + col_gap)
        ThreePillarCard.render(
            slide,
            x=cx,
            y=1.75,
            w=col_w,
            h=col_h,
            title=p[0],
            body=p[1],
            icon_path=p[2] if len(p) > 2 else None
        )

    # Supporting Artefact / Seal Far Right
    img = seal_image_path or get_image_asset("image-21-5.png")
    HistoricalImageFrame.render(
        slide,
        x=9.75,
        y=1.90,
        w=3.033,
        h=4.50,
        image_path=img,
        caption=caption or "Royal Sanskrit Rajmudra seal of Chhatrapati Shivaji",
        caption_h=0.45
    )


# ====================================================================
# M22 — DUAL HISTORICAL LEADERS
# ====================================================================
def render_m22_dual_leaders(
    slide: Any,
    kicker: str,
    title: str,
    leader1: Dict[str, str],  # {name, role, bio, image_path, caption}
    leader2: Dict[str, str],
    slide_num: int = 22
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=22)
    SlideFooter.render(slide, slide_num=slide_num)

    # Left Portrait
    img1 = leader1.get("image_path") or get_image_asset("image-22-2.png")
    HistoricalImageFrame.render(
        slide,
        x=0.55,
        y=1.75,
        w=2.85,
        h=4.80,
        image_path=img1,
        caption=leader1.get("caption", "Tarabai in battle"),
        caption_h=0.45
    )

    # Leader 1 Card
    DualBiographyCard.render(
        slide,
        x=3.65,
        y=1.75,
        w=3.05,
        h=4.80,
        name=leader1.get("name", "Tarabai"),
        role=leader1.get("role", "Warrior Queen"),
        bio=leader1.get("bio", "")
    )

    # Leader 2 Card
    DualBiographyCard.render(
        slide,
        x=6.95,
        y=1.75,
        w=3.05,
        h=4.80,
        name=leader2.get("name", "Ahilyabai Holkar"),
        role=leader2.get("role", "Philosopher Queen"),
        bio=leader2.get("bio", "")
    )

    # Right Portrait / Stamp
    img2 = leader2.get("image_path") or get_image_asset("image-22-3.png")
    HistoricalImageFrame.render(
        slide,
        x=10.25,
        y=1.75,
        w=2.533,
        h=4.80,
        image_path=img2,
        caption=leader2.get("caption", "Commemorative postage stamp"),
        caption_h=0.45
    )


# ====================================================================
# M23 — FEATURE STORY + THREE-IMAGE CULTURAL GALLERY
# ====================================================================
def render_m23_gallery_story(
    slide: Any,
    kicker: str,
    title: str,
    body_bullets: List[str],
    hero_image_path: Optional[str] = None,
    hero_caption: Optional[str] = None,
    sub_image1_path: Optional[str] = None,
    sub_caption1: Optional[str] = None,
    sub_image2_path: Optional[str] = None,
    sub_caption2: Optional[str] = None,
    slide_num: int = 23
):
    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=23)
    SlideFooter.render(slide, slide_num=slide_num)

    # Narrative Card Left
    ParchmentCard.render(
        slide,
        x=0.55,
        y=1.72,
        w=7.35,
        h=4.80,
        body_bullets=body_bullets,
        fill_color=COLORS.rgb_parchment,
        body_size_pt=13.5
    )

    # Hero Image Upper Right
    h_img = hero_image_path or get_image_asset("image-23-2.png")
    HistoricalImageFrame.render(
        slide,
        x=8.20,
        y=1.72,
        w=4.583,
        h=2.85,
        image_path=h_img,
        caption=hero_caption or "Thanjavur gold-leaf painting",
        caption_h=0.38
    )

    # Sub-image 1 Lower Right
    s1_img = sub_image1_path or get_image_asset("image-23-3.png")
    HistoricalImageFrame.render(
        slide,
        x=8.20,
        y=4.80,
        w=2.15,
        h=1.72,
        image_path=s1_img,
        caption=sub_caption1 or "Brihadishwara inscription",
        caption_h=0.36
    )

    # Sub-image 2 Lower Right
    s2_img = sub_image2_path or get_image_asset("image-23-4.png")
    HistoricalImageFrame.render(
        slide,
        x=10.633,
        y=4.80,
        w=2.15,
        h=1.72,
        image_path=s2_img,
        caption=sub_caption2 or "Modi script document",
        caption_h=0.36
    )


# ====================================================================
# M24 — DARK LEGACY SUMMARY
# ====================================================================
def render_m24_dark_summary(
    slide: Any,
    kicker: str,
    title: str,
    takeaways: List[str],
    closing_quote: Optional[str] = None,
    slide_num: int = 24
):
    set_slide_background(slide, COLORS.rgb_maroon)

    # Subtle decorative circular geometry (placed BEFORE header and content so it stays in the background)
    create_solid_shape(slide, MSO_SHAPE.OVAL, -1.1, -1.1, 3.0, 3.0, fill_color=COLORS.rgb_maroon_alt)
    create_solid_shape(slide, MSO_SHAPE.OVAL, 11.3, 5.3, 3.2, 3.2, fill_color=COLORS.rgb_maroon_alt)

    MarathaEditorialHeader.render(slide, kicker=kicker, title=title, slide_idx=24, is_dark=True)
    SlideFooter.render(slide, slide_num=slide_num, is_dark=True)

    MarathaLegacySummary.render(
        slide,
        x=0.60,
        y=1.85,
        w=11.60,
        h=4.80,
        takeaways=takeaways,
        final_closing_text=closing_quote or "Driven throughout by the fiery ideal of Swarajya — self-rule grounded in justice, security, and cultural dignity."
    )
