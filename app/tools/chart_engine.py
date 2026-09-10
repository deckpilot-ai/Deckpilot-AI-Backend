"""Native PowerPoint chart generation and enterprise styling engine."""

import logging
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from app.schemas.generation_state import ChartSpec, ChartType, DesignSystem

logger = logging.getLogger(__name__)


def hex_to_rgb(hex_str: str) -> RGBColor:
    val = hex_str.lstrip("#")
    try:
        return RGBColor(int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16))
    except Exception:
        return RGBColor(37, 99, 235)


class ChartEngine:
    """Renders native, editable Microsoft PowerPoint charts with enterprise styling."""

    CHART_TYPE_MAP = {
        ChartType.COLUMN: XL_CHART_TYPE.COLUMN_CLUSTERED,
        ChartType.BAR: XL_CHART_TYPE.BAR_CLUSTERED,
        ChartType.STACKED_COLUMN: XL_CHART_TYPE.COLUMN_STACKED,
        ChartType.STACKED_BAR: XL_CHART_TYPE.BAR_STACKED,
        ChartType.LINE: XL_CHART_TYPE.LINE,
        ChartType.AREA: XL_CHART_TYPE.AREA,
        ChartType.PIE: XL_CHART_TYPE.PIE,
        ChartType.DONUT: XL_CHART_TYPE.DOUGHNUT,
    }

    @classmethod
    def render_chart(
        cls,
        slide: Any,
        chart_spec: ChartSpec,
        design_system: DesignSystem,
        x: float,
        y: float,
        w: float,
        h: float,
    ) -> Any:
        """Adds a native OpenXML chart shape to a slide."""
        if not chart_spec.categories or not chart_spec.series:
            logger.warning("Empty chart specification skipped")
            return None

        chart_data = CategoryChartData()
        chart_data.categories = chart_spec.categories

        for s in chart_spec.series:
            name = s.get("name", "Metric")
            raw_vals = s.get("values", [])
            # Coerce numbers safely
            clean_vals = []
            for v in raw_vals:
                try:
                    clean_vals.append(float(v))
                except (ValueError, TypeError):
                    clean_vals.append(0.0)
            chart_data.add_series(name, clean_vals)

        xl_type = cls.CHART_TYPE_MAP.get(chart_spec.chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)

        shape = slide.shapes.add_chart(
            xl_type,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
            chart_data,
        )
        chart = shape.chart
        chart.has_legend = len(chart_spec.series) > 1 or chart_spec.chart_type in (ChartType.PIE, ChartType.DONUT)
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            if chart.legend.font:
                chart.legend.font.name = design_system.typography.body_font.name
                chart.legend.font.size = Pt(11)

        # Style chart title if specified
        if chart_spec.title:
            chart.has_title = True
            chart.chart_title.text_frame.text = chart_spec.title
            if chart.chart_title.text_frame.paragraphs:
                p = chart.chart_title.text_frame.paragraphs[0]
                p.font.name = design_system.typography.title_font.name
                p.font.size = Pt(13)
                p.font.bold = True
                p.font.color.rgb = hex_to_rgb(design_system.colors.text_primary)
        else:
            chart.has_title = False

        # Apply series palette colors
        palette = [hex_to_rgb(c) for c in design_system.colors.chart_colors]
        try:
            for idx, series in enumerate(chart.series):
                color = palette[idx % len(palette)]
                if hasattr(series, "format") and hasattr(series.format, "fill"):
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = color
                if hasattr(series, "format") and hasattr(series.format, "line"):
                    if chart_spec.chart_type == ChartType.LINE:
                        series.format.line.color.rgb = color
                        series.format.line.width = Pt(2.5)
                    else:
                        series.format.line.fill.background()
        except Exception as e:
            logger.warning("Chart series color formatting note: %s", e)

        return shape
