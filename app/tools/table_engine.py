"""Native enterprise table generator tool with presentation-ready styling."""

import logging
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.schemas.generation_state import DesignSystem, TableSpec

logger = logging.getLogger(__name__)


def hex_to_rgb(hex_str: str, default: tuple[int, int, int] = (19, 42, 82)) -> RGBColor:
    val = hex_str.lstrip("#")
    try:
        return RGBColor(int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16))
    except Exception:
        return RGBColor(*default)


def tint_rgb(color: RGBColor, amount: float) -> RGBColor:
    return RGBColor(*(round(c + (255 - c) * amount) for c in color))


class TableEngine:
    """Renders native Microsoft PowerPoint structured tables with enterprise styling."""

    @classmethod
    def render_table(
        cls,
        slide: Any,
        table_spec: TableSpec,
        design_system: DesignSystem,
        x: float,
        y: float,
        w: float,
        h: float,
    ) -> Any:
        if not table_spec.headers and not table_spec.rows:
            return None

        num_cols = max(len(table_spec.headers), max((len(r) for r in table_spec.rows), default=1))
        num_rows = (1 if table_spec.headers else 0) + len(table_spec.rows)

        shape = slide.shapes.add_table(num_rows, num_cols, Inches(x), Inches(y), Inches(w), Inches(h))
        table = shape.table

        # Colors
        primary = hex_to_rgb(design_system.colors.primary)
        accent = hex_to_rgb(design_system.colors.accent)
        zebra_fill = tint_rgb(primary, 0.95)
        white_fill = RGBColor(255, 255, 255)
        text_primary = hex_to_rgb(design_system.colors.text_primary)
        font_name = design_system.typography.body_font.name

        row_offset = 0

        # Render Header Row
        if table_spec.headers:
            for col_idx, header_text in enumerate(table_spec.headers):
                cell = table.cell(0, col_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = primary
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE

                # Text formatting
                p = cell.text_frame.paragraphs[0]
                p.text = str(header_text)
                p.font.name = font_name
                p.font.size = Pt(11.5)
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = PP_ALIGN.CENTER
            row_offset = 1

        # Render Data Rows
        for r_idx, row_data in enumerate(table_spec.rows):
            current_row = r_idx + row_offset
            is_zebra = table_spec.zebra_striping and (r_idx % 2 == 1)
            row_fill = zebra_fill if is_zebra else white_fill

            for col_idx in range(num_cols):
                cell = table.cell(current_row, col_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_fill
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE

                cell_value = str(row_data[col_idx]) if col_idx < len(row_data) else ""
                p = cell.text_frame.paragraphs[0]
                p.text = cell_value
                p.font.name = font_name
                p.font.size = Pt(10.5)
                p.font.color.rgb = text_primary

                # Align right if numeric/currency/percent
                if any(c.isdigit() for c in cell_value) and (cell_value.endswith("%") or cell_value.startswith("$") or cell_value.endswith("x")):
                    p.alignment = PP_ALIGN.RIGHT
                elif col_idx == 0:
                    p.alignment = PP_ALIGN.LEFT
                    p.font.bold = True
                else:
                    p.alignment = PP_ALIGN.LEFT

        return shape
