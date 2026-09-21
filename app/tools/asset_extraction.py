"""Document and asset extraction tool for multi-format inputs."""

import csv
import hashlib
import io
import logging
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
import pymupdf
from docx import Document
from PIL import Image, ImageOps
from pptx import Presentation

from app.schemas.generation_state import AssetMetadata

logger = logging.getLogger(__name__)


class ExtractedAssetPayload:
    def __init__(self, storage_key: str, data: bytes, content_type: str, metadata: AssetMetadata) -> None:
        self.storage_key = storage_key
        self.data = data
        self.content_type = content_type
        self.metadata = metadata


class ExtractionOutput:
    def __init__(self) -> None:
        self.text_blocks: list[dict[str, Any]] = []
        self.tables: list[dict[str, Any]] = []
        self.assets: list[ExtractedAssetPayload] = []
        self.metadata: dict[str, Any] = {}


class DocumentAssetExtractor:
    """Deterministic extractor for extracting text, tables, and visual assets from documents."""

    @classmethod
    def extract_document(
        cls,
        file_bytes: bytes,
        filename: str,
        mime_type: str = "",
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        ext = Path(filename).suffix.lower()
        if ext in [".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tiff"]:
            return cls.extract_image_file(file_bytes, filename, mime_type, on_progress)
        elif ext == ".pdf":
            return cls.extract_pdf(file_bytes, filename, on_progress)
        elif ext in [".docx", ".doc"]:
            return cls.extract_docx(file_bytes, filename, on_progress)
        elif ext in [".xlsx", ".xlsm", ".xls"]:
            return cls.extract_spreadsheet(file_bytes, filename, on_progress)
        elif ext in [".csv", ".tsv"]:
            return cls.extract_csv(file_bytes, filename, on_progress)
        elif ext == ".pptx":
            return cls.extract_pptx(file_bytes, filename, on_progress)
        elif ext in [".txt", ".md", ".json", ".yaml", ".yml", ".xml"]:
            return cls.extract_text(file_bytes, filename, on_progress)
        else:
            return cls.extract_text(file_bytes, filename, on_progress)

    @classmethod
    def extract_image_file(
        cls,
        file_bytes: bytes,
        filename: str,
        mime_type: str = "image/png",
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        output = ExtractionOutput()
        sha256 = hashlib.sha256(file_bytes).hexdigest()
        ext = Path(filename).suffix.lstrip(".").lower() or "png"
        if ext == "jpeg":
            ext = "jpg"

        try:
            with Image.open(io.BytesIO(file_bytes)) as img:
                normalized = ImageOps.exif_transpose(img).convert("RGBA" if "A" in img.getbands() else "RGB")
                width, height = normalized.size
                aspect_ratio = round(width / max(1, height), 3)
                
                # Encode normalized PNG
                buffer = io.BytesIO()
                normalized.save(buffer, format="PNG")
                png_bytes = buffer.getvalue()
        except Exception as exc:
            logger.warning("Failed to decode image file %s: %s", filename, exc)
            return output

        storage_key = f"extracted/images/{sha256[:16]}_{Path(filename).stem}.png"
        meta = AssetMetadata(
            asset_id=f"img_{sha256[:12]}",
            source_file=filename,
            page_or_slide=1,
            width=width,
            height=height,
            aspect_ratio=aspect_ratio,
            format="png",
            sha256=sha256,
            caption=Path(filename).stem.replace("_", " ").replace("-", " ").title(),
            nearby_text=f"Uploaded image asset: {filename}",
            storage_key=storage_key,
            quality_score=1.0,
            is_valid_figure=True,
        )

        output.assets.append(ExtractedAssetPayload(storage_key, png_bytes, "image/png", meta))
        output.text_blocks.append({
            "content": f"### Uploaded Image Asset: {filename}\n- Dimensions: {width}x{height} (Aspect Ratio: {aspect_ratio})\n- Context: {meta.caption}",
            "source": filename,
        })
        return output

    @classmethod
    def extract_pdf(
        cls,
        file_bytes: bytes,
        filename: str,
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        from app.services.extraction import DocumentExtractor

        extract_res = DocumentExtractor.extract_pdf(file_bytes, filename, on_progress=on_progress)

        output = ExtractionOutput()
        output.text_blocks = extract_res.text_blocks
        output.tables = extract_res.tables
        output.metadata = extract_res.metadata

        # Build ExtractedAssetPayload with rich AssetMetadata for each extracted visual figure
        payload_map = {key: (data, ct) for key, data, ct in extract_res.image_payloads}
        for img in extract_res.extracted_images:
            storage_key = img.get("storage_key", "")
            data, content_type = payload_map.get(storage_key, (b"", "image/png"))
            if not data:
                continue
            meta = AssetMetadata(
                asset_id=img.get("asset_id") or f"art_{img.get('sha256', '')[:12]}",
                source_file=filename,
                page_or_slide=img.get("source_page", 1),
                width=img.get("width", 800),
                height=img.get("height", 600),
                aspect_ratio=img.get("aspect_ratio", 1.33),
                format=storage_key.split(".")[-1] if "." in storage_key else "png",
                sha256=img.get("sha256", ""),
                image_type=img.get("image_type", "photograph"),
                section=img.get("section", ""),
                semantic_tags=img.get("semantic_tags", []),
                caption=img.get("caption", ""),
                nearby_text=img.get("nearby_text", ""),
                storage_key=storage_key,
                quality_score=img.get("quality_score", 0.95),
                is_valid_figure=True,
            )
            output.assets.append(ExtractedAssetPayload(storage_key, data, content_type, meta))

        return output

    @classmethod
    def extract_docx(
        cls,
        file_bytes: bytes,
        filename: str,
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        output = ExtractionOutput()
        doc = Document(io.BytesIO(file_bytes))

        full_text = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                if p.style and "Heading" in p.style.name:
                    full_text.append(f"### {t}")
                else:
                    full_text.append(t)

        if full_text:
            output.text_blocks.append({
                "content": "\n\n".join(full_text),
                "source": filename,
            })

        for table_idx, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            if rows:
                output.tables.append({
                    "table_index": table_idx + 1,
                    "rows": rows,
                    "source": f"{filename}#table={table_idx + 1}",
                })
                # Markdown representation
                md_table = [f"| {' | '.join(rows[0])} |", f"| {' | '.join(['---'] * len(rows[0]))} |"]
                for r in rows[1:]:
                    md_table.append(f"| {' | '.join(r)} |")
                output.text_blocks.append({
                    "content": f"### Document Table {table_idx + 1}:\n" + "\n".join(md_table),
                    "source": f"{filename}#table={table_idx + 1}",
                })

        try:
            for rel_id, part in doc.part.related_parts.items():
                if "image" in part.content_type:
                    img_bytes = part.blob
                    ext = part.content_type.split("/")[-1]
                    sha256 = hashlib.sha256(img_bytes).hexdigest()
                    storage_key = f"extracted/{Path(filename).stem}_{rel_id}.{ext}"

                    with Image.open(io.BytesIO(img_bytes)) as img:
                        w, h = img.size
                        aspect = round(w / max(1, h), 3)

                    meta = AssetMetadata(
                        asset_id=f"art_{sha256[:12]}",
                        source_file=filename,
                        page_or_slide=1,
                        width=w,
                        height=h,
                        aspect_ratio=aspect,
                        format=ext,
                        sha256=sha256,
                        caption=f"Embedded graphic from {filename}",
                        storage_key=storage_key,
                        quality_score=0.9,
                        is_valid_figure=True,
                    )
                    output.assets.append(ExtractedAssetPayload(storage_key, img_bytes, part.content_type, meta))
        except Exception as e:
            logger.warning("DOCX image extraction advisory: %s", e)

        return output

    @classmethod
    def extract_pptx(
        cls,
        file_bytes: bytes,
        filename: str,
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        output = ExtractionOutput()
        prs = Presentation(io.BytesIO(file_bytes))
        output.metadata["slide_count"] = len(prs.slides)
        output.metadata["slide_width"] = prs.slide_width
        output.metadata["slide_height"] = prs.slide_height

        for slide_idx, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    slide_texts.append(shape.text.strip())
                if hasattr(shape, "image"):
                    try:
                        img = shape.image
                        img_bytes = img.blob
                        ext = img.ext or "png"
                        sha256 = hashlib.sha256(img_bytes).hexdigest()
                        storage_key = f"extracted/{Path(filename).stem}_s{slide_idx+1}_{sha256[:8]}.{ext}"
                        w, h = shape.width.inches, shape.height.inches
                        aspect = round(w / max(0.01, h), 3)
                        meta = AssetMetadata(
                            asset_id=f"ppt_img_{sha256[:12]}",
                            source_file=filename,
                            page_or_slide=slide_idx + 1,
                            width=int(w * 100),
                            height=int(h * 100),
                            aspect_ratio=aspect,
                            format=ext,
                            sha256=sha256,
                            caption=f"Slide {slide_idx+1} visual asset",
                            storage_key=storage_key,
                            quality_score=0.95,
                            is_valid_figure=True,
                        )
                        output.assets.append(ExtractedAssetPayload(storage_key, img_bytes, f"image/{ext}", meta))
                    except Exception as e:
                        logger.warning("PPTX image extraction skip: %s", e)

            if slide_texts:
                output.text_blocks.append({
                    "slide": slide_idx + 1,
                    "content": "\n".join(slide_texts),
                    "source": f"{filename}#slide={slide_idx + 1}",
                })

        return output

    @classmethod
    def extract_spreadsheet(
        cls,
        file_bytes: bytes,
        filename: str,
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        output = ExtractionOutput()
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        output.metadata["sheet_names"] = wb.sheetnames

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append([str(c).strip() if c is not None else "" for c in row])

            if rows:
                output.tables.append({
                    "sheet": sheet_name,
                    "rows": rows[:150],
                    "source": f"{filename}#{sheet_name}",
                })
                md_lines = [f"### Spreadsheet Sheet: {sheet_name}"]
                if rows:
                    header = " | ".join(rows[0])
                    sep = " | ".join(["---"] * len(rows[0]))
                    md_lines.append(f"| {header} |")
                    md_lines.append(f"| {sep} |")
                    for r in rows[1:60]:
                        md_lines.append(f"| {' | '.join(r)} |")
                output.text_blocks.append({
                    "content": "\n".join(md_lines),
                    "source": f"{filename}#{sheet_name}",
                })
        return output

    @classmethod
    def extract_csv(
        cls,
        file_bytes: bytes,
        filename: str,
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        output = ExtractionOutput()
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")

        delim = "\t" if filename.endswith(".tsv") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delim)
        rows = [row for row in reader if any(cell.strip() for cell in row)]

        if rows:
            output.tables.append({"rows": rows[:150], "source": filename})
            md_lines = [f"### CSV Data Table: {filename}"]
            header = " | ".join(rows[0])
            sep = " | ".join(["---"] * len(rows[0]))
            md_lines.append(f"| {header} |")
            md_lines.append(f"| {sep} |")
            for r in rows[1:60]:
                md_lines.append(f"| {' | '.join(r)} |")
            output.text_blocks.append({"content": "\n".join(md_lines), "source": filename})
        return output

    @classmethod
    def extract_text(
        cls,
        file_bytes: bytes,
        filename: str,
        on_progress: Callable[[str], None] | None = None,
    ) -> ExtractionOutput:
        output = ExtractionOutput()
        try:
            content = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            content = file_bytes.decode("latin-1", errors="ignore")

        if content.strip():
            output.text_blocks.append({"content": content.strip(), "source": filename})
        return output
