"""Regression coverage for source-backed consulting decks and the live PDF failure."""

import io

import pytest
from pptx import Presentation
from sqlalchemy.pool import NullPool

from app.core.config import settings
from app.db.engine import _create_engine
from app.services.design_system import default_brand, prepare_deck
from app.services.prompts import (
    ASK_MODE_SYSTEM_PROMPT,
    DECK_PLANNER_SYSTEM_PROMPT,
    PLAN_MODE_SYSTEM_PROMPT,
    SLIDE_WRITER_SYSTEM_PROMPT,
)
from app.services.renderer import PPTXRenderer


def text_of(payload):
    prs = Presentation(io.BytesIO(payload))
    return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)


def test_design_injection_is_deck_specific():
    for prompt in (PLAN_MODE_SYSTEM_PROMPT, DECK_PLANNER_SYSTEM_PROMPT, SLIDE_WRITER_SYSTEM_PROMPT):
        assert "CONSULTING-PPTX-DESIGN v1.0" in prompt
        assert "No em dashes" in prompt
    assert "CONSULTING-PPTX-DESIGN" not in ASK_MODE_SYSTEM_PROMPT
    assert "Include realistic, high-fidelity numbers" not in SLIDE_WRITER_SYSTEM_PROMPT


def test_stats_and_quotes_require_evidence():
    source = 'Retention reached 98%. Researcher Rao wrote: "The script remains undeciphered."'
    spec = {"slides": [{"layoutHint": "metrics_grid", "metrics": [
        {"value": "98%", "label": "Retention", "sourceQuote": "Retention reached 98%."},
        {"value": "140%", "label": "Growth", "sourceQuote": "Growth reached 140%."},
    ]}, {"layoutHint": "quote", "quote": {
        "text": "The script remains undeciphered.", "attribution": "Rao", "sourceQuote": source,
    }}, {"layoutHint": "quote", "quote": {
        "text": "A fabricated quote", "attribution": "Rao", "sourceQuote": source,
    }}]}
    result = prepare_deck(spec, source)
    assert len(result["slides"][0]["metrics"]) == 1
    assert result["slides"][1]["quote"]["attribution"] == "Rao"
    assert "quote" not in result["slides"][2]


def test_render_contains_source_content_without_sample_statistics():
    spec = prepare_deck({"deckTitle": "Harappan History", "slides": [
        {"purpose": "Chronology", "message": "Harappan phases", "layoutHint": "timeline",
         "bullets": ["Early phase - 6000 BCE to 2600 BCE", "Mature phase - 2600 BCE to 1900 BCE", "Late phase - 1900 BCE to 1300 BCE"]},
        {"purpose": "Evidence", "message": "Evidence remains open to interpretation", "layoutHint": "metrics_grid", "bullets": ["The script remains undeciphered."]},
    ]}, "")
    payload = PPTXRenderer.render_deck(spec)
    text = text_of(payload)
    for invented in ("+140%", "99.4%", "3.8x", "Discovery & Ingestion", "Validated assumptions"):
        assert invented not in text
    assert "6000 BCE" in text
    assert "script remains undeciphered" in text
    assert PPTXRenderer.validate_deck(payload, 2)["visualInspection"] == "not_performed"


def test_dense_numbered_cards_preserve_text():
    bullets = [f"Evidence item {i}: Archaeologists distinguish observed remains from uncertain historical interpretations." for i in range(6)]
    payload = PPTXRenderer.render_deck({"deckTitle": "History", "slides": [{
        "purpose": "Interpretation", "message": "Evidence requires careful interpretation",
        "layoutHint": "roadmap", "bullets": bullets,
    }]})
    for bullet in bullets:
        assert bullet in text_of(payload)
    PPTXRenderer.validate_deck(payload, 1)


def test_oversized_copy_fails_before_delivery():
    with pytest.raises(ValueError, match="layout budget"):
        PPTXRenderer.render_deck({"slides": [{"message": "Long title " * 100}]})


def test_dense_cards_reflow_without_losing_source_copy():
    bullets = [f"Evidence {i}: " + "Archaeologists compare domestic architecture, drainage and craft remains to interpret social organisation cautiously. " * 2 for i in range(6)]
    payload = PPTXRenderer.render_deck({"deckTitle": "History", "slides": [{
        "message": "Reading archaeological evidence", "layoutHint": "roadmap", "bullets": bullets,
    }]})
    for bullet in bullets:
        assert bullet.strip() in text_of(payload)
    PPTXRenderer.validate_deck(payload, 1)


def test_cover_preserves_all_six_bullets_with_takeaway():
    bullets = [f"Source item {i}: Material remains reveal urban planning, trade and craft production across the region." for i in range(6)]
    payload = PPTXRenderer.render_deck({"deckTitle": "History", "slides": [{
        "message": "The Harappan Civilisation", "layoutHint": "hero", "bullets": bullets,
        "takeaway": "Evidence supports multiple interpretations.",
    }]})
    for bullet in bullets:
        assert bullet in text_of(payload)


def test_fonts_palette_and_em_dash_normalization():
    assert default_brand("Indian history")["colors"] != default_brand("Climate sustainability")["colors"]
    spec = prepare_deck({"deckTitle": "History", "slides": [{
        "purpose": "History", "message": "Bricks — beads", "layoutHint": "hero",
        "bullets": ["Evidence from the source"], "speakerNotes": "Page 1 — source",
    }]}, "")
    payload = PPTXRenderer.render_deck(spec)
    assert "\u2014" not in text_of(payload)
    prs = Presentation(io.BytesIO(payload))
    names = {p.font.name for shape in prs.slides[0].shapes if shape.has_text_frame for p in shape.text_frame.paragraphs}
    assert {"Cambria", "Calibri"} <= names
    assert "Page 1 - source" in prs.slides[0].notes_slide.notes_text_frame.text


def test_turso_sessions_do_not_share_singleton_thread_connections(monkeypatch):
    monkeypatch.setattr(settings, "turso_database_url", "libsql://example.invalid")
    monkeypatch.setattr(settings, "turso_auth_token", "test-token")
    engine = _create_engine()
    assert isinstance(engine.pool, NullPool)
    engine.dispose()


def test_content_boxes_use_available_canvas_and_subjects_have_distinct_covers():
    outputs = []
    for title in ['The Rise of Empires', 'The Parliamentary System', 'Factors of Production']:
        spec = prepare_deck({'deckTitle': title, 'slides': [
            {'message': title, 'bullets': ['A chapter study guide.']},
            {'message': 'Compare the evidence', 'layoutHint': 'comparison', 'bullets': ['First point.', 'Second point.']},
            {'message': 'What should we remember?', 'bullets': ['Discuss the evidence and its implications.']},
        ]}, '')
        assert spec['slides'][0]['layoutHint'] == 'hero'
        assert spec['slides'][-1]['layoutHint'] == 'closing'
        outputs.append(PPTXRenderer.render_deck(spec))
    decks = [Presentation(io.BytesIO(p)) for p in outputs]
    assert len({str(p.slides[0].background.fill.fore_color.rgb) for p in decks}) == 3
    for p in decks:
        cards = [s for s in p.slides[1].shapes if s.name == 'card']
        assert cards and all(s.height / 914400 >= 3.8 for s in cards)


def test_qr_assets_are_rejected_and_real_images_are_accepted():
    import cv2
    import numpy as np

    from app.services.image_quality import is_documentary_image
    qr = cv2.QRCodeEncoder_create().encode('https://example.invalid/textbook')
    qr = cv2.copyMakeBorder(qr, 4, 4, 4, 4, cv2.BORDER_CONSTANT, value=255)
    qr = cv2.resize(qr, (300, 300), interpolation=cv2.INTER_NEAREST)
    assert not is_documentary_image(cv2.imencode('.png', qr)[1].tobytes())
    photo = np.full((120, 200, 3), (110, 170, 220), dtype=np.uint8)
    cv2.circle(photo, (90, 60), 30, (30, 80, 120), -1)
    assert is_documentary_image(cv2.imencode('.png', photo)[1].tobytes())
    assert not is_documentary_image(b'invalid image')


def test_source_backed_chart_uses_real_comparable_values():
    evidence = 'Land 20. Labour 40.'
    spec = prepare_deck({'slides': [{'message': 'Source comparison', 'layoutHint': 'bar_chart', 'metrics': [
        {'value': '20', 'label': 'Land', 'sourceQuote': 'Land 20.'},
        {'value': '40', 'label': 'Labour', 'sourceQuote': 'Labour 40.'},
    ]}]}, evidence)
    payload = PPTXRenderer.render_deck(spec)
    prs = Presentation(io.BytesIO(payload))
    bars = [s for s in prs.slides[0].shapes if s.name == 'data-bar']
    assert len(bars) == 2 and bars[1].width == 2 * bars[0].width
    assert '20' in text_of(payload) and '40' in text_of(payload)


def test_image_slide_and_grouped_cards_use_native_bullets_with_hanging_indent():
    from PIL import Image

    photo = io.BytesIO()
    Image.new('RGB', (100, 100), 'orange').save(photo, format='PNG')
    points = [('A long first point that wraps across lines while retaining its own bullet marker. ' * 2).strip(),
              'Second point: Separate explanation.', 'A third distinct point.']
    spec = {'slides': [
        {'message': 'Image evidence', 'layoutHint': 'two_column', 'bullets': points, 'imageArtifactId': 'image'},
        {'message': 'Grouped concepts', 'layoutHint': 'comparison', 'bullets': points},
    ]}
    prs = Presentation(io.BytesIO(PPTXRenderer.render_deck(spec, source_images={'image': photo.getvalue()})))
    for slide in prs.slides:
        box = next(s for s in slide.shapes if s.has_text_frame and points[0] in s.text)
        paragraphs = box.text_frame.paragraphs
        assert paragraphs[0].text == points[0]
        assert paragraphs[1].text == points[1]
        for p in paragraphs:
            properties = p._p.get_or_add_pPr()
            assert properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar').get('char') == '•'
            assert int(properties.get('marL')) > 0
            assert int(properties.get('indent')) < 0
        assert paragraphs[0].space_after.pt == 5
        assert paragraphs[0].line_spacing == 1.15
        title = next(s for s in slide.shapes if s.has_text_frame and s.text in ('Image evidence', 'Grouped concepts'))
        assert not title.text_frame.paragraphs[0]._p.xpath('./a:pPr/a:buChar')
