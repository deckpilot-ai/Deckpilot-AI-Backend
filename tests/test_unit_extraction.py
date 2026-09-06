"""Unit tests for document and multi-modal asset extraction service and agents."""

import io

import fitz  # PyMuPDF
import openpyxl
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from app.agents.extraction_agent import (
    ImageExtractionAgent,
    ReferenceIntakeAgent,
    TextExtractionAgent,
)
from app.services.extraction import DocumentExtractor


def test_extract_text_and_markdown():
    text_data = b"# Executive Summary\n\ndeckpilotAI delivers presentations from prompts and reference documents."
    res = DocumentExtractor.extract_document(text_data, "summary.md", "text/markdown")
    assert len(res.text_blocks) == 1
    assert "deckpilotAI" in res.text_blocks[0]["content"]


def test_extract_csv_spreadsheet():
    csv_data = b"Quarter,Revenue,Growth\nQ1,1.2M,15%\nQ2,1.8M,50%\nQ3,2.7M,50%"
    res = DocumentExtractor.extract_document(csv_data, "financials.csv", "text/csv")
    assert len(res.tables) == 1
    assert len(res.text_blocks) == 1
    assert "Revenue" in res.text_blocks[0]["content"]


def test_extract_pptx_slides():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Existing Sample Slide"
    slide.placeholders[1].text = "Existing sample subtitle notes"

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    res = DocumentExtractor.extract_document(
        pptx_bytes,
        "sample.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert res.metadata["slide_count"] == 1
    assert len(res.text_blocks) == 1
    assert "Existing Sample Slide" in res.text_blocks[0]["content"]


def test_extract_image_jpg():
    """Verify JPG/PNG image extraction extracts dimensions, saves asset, and generates visual metadata."""
    img = Image.new("RGB", (1280, 720), color=(10, 20, 30))
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    jpg_bytes = buf.getvalue()

    res = DocumentExtractor.extract_document(jpg_bytes, "architecture_diagram.jpg", "image/jpeg")

    assert res.metadata["width"] == 1280
    assert res.metadata["height"] == 720
    assert res.metadata["aspect_ratio"] == 1.78
    assert len(res.extracted_images) == 1
    assert res.extracted_images[0]["format"] == "png"
    assert res.image_payloads[0][2] == "image/png"
    assert "extracted/images/" in res.extracted_images[0]["storage_key"]

    # Text block metadata
    assert len(res.text_blocks) == 1
    assert "architecture_diagram.jpg" in res.text_blocks[0]["content"]
    assert "1280x720" in res.text_blocks[0]["content"]


def test_extract_word_docx():
    """Verify Microsoft Word .docx extraction parses paragraphs and tables."""
    doc = Document()
    doc.add_heading("Product Strategy 2026", level=1)
    doc.add_paragraph("deckpilotAI transforms ideas into executive PowerPoint slides.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Multi-agent engine"
    table.cell(1, 1).text = "Operational"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = DocumentExtractor.extract_document(
        docx_bytes,
        "strategy.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert len(res.text_blocks) >= 2  # Heading/paragraphs + Markdown table
    assert any("Product Strategy 2026" in tb["content"] for tb in res.text_blocks)
    assert len(res.tables) == 1
    assert res.tables[0]["rows"][0] == ["Category", "Status"]
    assert res.tables[0]["rows"][1] == ["Multi-agent engine", "Operational"]


def test_extract_excel_xlsx():
    """Verify Microsoft Excel .xlsx extraction parses workbooks, sheets, and rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.append(["Year", "ARR", "Net Retention"])
    ws.append(["2024", "$1.2M", "118%"])
    ws.append(["2025", "$3.4M", "135%"])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    res = DocumentExtractor.extract_document(
        xlsx_bytes,
        "metrics.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert "Financials" in res.metadata["sheet_names"]
    assert len(res.tables) == 1
    assert res.tables[0]["sheet"] == "Financials"
    assert res.tables[0]["rows"][0] == ["Year", "ARR", "Net Retention"]
    assert res.tables[0]["rows"][1] == ["2024", "$1.2M", "118%"]


def test_extract_pdf_document():
    """Verify PDF extraction parses pages and text."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Executive Brief: Q3 Board Review\nRevenue grew by 45% MoM.")

    pdf_bytes = doc.tobytes()
    doc.close()

    res = DocumentExtractor.extract_document(pdf_bytes, "brief.pdf", "application/pdf")
    assert res.metadata["page_count"] == 1
    assert len(res.text_blocks) == 1
    assert "Executive Brief: Q3 Board Review" in res.text_blocks[0]["content"]


@pytest.mark.asyncio
async def test_extraction_agents_integration():
    """Verify ReferenceIntakeAgent, TextExtractionAgent, and ImageExtractionAgent execute properly."""
    # 1. Text Agent on CSV
    text_agent = TextExtractionAgent()
    csv_bytes = b"Product,Units,Revenue\nAlpha,100,5000\nBeta,200,12000"
    text_res = await text_agent.run({
        "file_bytes": csv_bytes,
        "filename": "sales.csv",
        "mime_type": "text/csv",
    })
    assert text_res["status"] == "success"
    assert len(text_res["tables"]) == 1

    # 2. Image Agent on PNG
    img = Image.new("RGB", (640, 480), color="red")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    png_bytes = buf.getvalue()

    img_agent = ImageExtractionAgent()
    img_res = await img_agent.run({
        "file_bytes": png_bytes,
        "filename": "mockup.png",
        "mime_type": "image/png",
    })
    assert img_res["status"] == "success"
    assert len(img_res["extracted_images"]) == 1
    assert img_res["metadata"]["width"] == 640
    assert img_res["metadata"]["height"] == 480

    # 3. Reference Intake Agent
    intake_agent = ReferenceIntakeAgent()
    intake_res = await intake_agent.run({
        "file_bytes": png_bytes,
        "filename": "mockup.png",
        "mime_type": "image/png",
    })
    assert intake_res["status"] == "success"
    assert intake_res["images_count"] == 1
