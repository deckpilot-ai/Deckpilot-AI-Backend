import io

from pptx import Presentation
from pptx.util import Inches, Pt

from app.agents.qa_agent import PresentationQAAgent
from app.agents.qa_checkpoints import CHECKPOINTS, checkpoint_ids
from app.agents.repair_agent import RepairAgent
from app.schemas.generation_state import (
    AssetMetadata,
    DesignSystem,
    LayoutFamily,
    QAReport,
    SlideSpec,
    ValidationCategory,
    ValidationIssue,
    ValidationSeverity,
)
from app.services.image_matcher import ImageMatcher
from app.services.renderer import PPTXRenderer


def test_checkpoint_catalog_has_120_stable_unique_checks():
    ids = checkpoint_ids()
    assert len(CHECKPOINTS) == 120
    assert len(ids) == len(set(ids))
    assert ids[0] == "QA-001"
    assert ids[-1] == "QA-120"
    assert all(item.repair_action for item in CHECKPOINTS)


def test_qa_detects_and_repairs_duplicate_copy_and_images():
    slides = [
        SlideSpec(slide_id="s01", slide_number=1, headline="Solar Energy", layout_family=LayoutFamily.HERO),
        SlideSpec(
            slide_id="s02",
            slide_number=2,
            headline="Solar Panel Economics",
            bullets=["Lower costs improve adoption", "Lower costs improve adoption"],
            image_artifact_id="solar-1",
            image_caption="Solar panels on a commercial rooftop",
            layout_family=LayoutFamily.TEXT_IMAGE,
        ),
        SlideSpec(
            slide_id="s03",
            slide_number=3,
            headline="Wind Farm Operations",
            bullets=["Turbine maintenance improves availability"],
            image_artifact_id="solar-1",
            image_caption="Solar panels on a commercial rooftop",
            layout_family=LayoutFamily.TEXT_IMAGE,
        ),
    ]
    assets = [
        AssetMetadata(asset_id="solar-1", caption="Solar panels on a commercial rooftop"),
        AssetMetadata(asset_id="wind-1", caption="Wind turbines at an operating wind farm"),
    ]

    report = PresentationQAAgent.evaluate_presentation(slides, DesignSystem(), available_assets=assets)
    failed = {issue.checkpoint_id for issue in report.issues}
    assert "QA-009" in failed
    assert "QA-061" in failed
    assert "QA-062" in failed
    assert report.checkpoints_total == 120
    assert report.repair_triggered is True

    repaired = RepairAgent.apply_corrections(slides, report, available_assets=assets)
    assert repaired[1].bullets == ["Lower costs improve adoption"]
    assert repaired[2].image_artifact_id == "wind-1"
    assert len({s.image_artifact_id for s in repaired if s.image_artifact_id}) == 2


def test_rendered_small_font_and_blank_space_are_detected_and_scaled():
    spec = SlideSpec(
        slide_id="s01",
        slide_number=1,
        headline="Readable typography",
        bullets=["A short but meaningful supporting point"],
        layout_family=LayoutFamily.TWO_COLUMN,
    )
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(2.0), Inches(0.4))
    box.text_frame.paragraphs[0].add_run().text = "Tiny body copy"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    page = slide.shapes.add_textbox(Inches(12.1), Inches(6.9), Inches(0.5), Inches(0.3))
    page.text = "1"
    stream = io.BytesIO()
    prs.save(stream)

    report = PresentationQAAgent.evaluate_presentation([spec], DesignSystem(), pptx_bytes=stream.getvalue())
    failed = {issue.checkpoint_id for issue in report.issues}
    assert "QA-025" in failed
    assert "QA-049" in failed

    repaired = RepairAgent.apply_corrections([spec], report)
    assert repaired[0].archetype_fields["qa_font_scale"] >= 1.15
    assert repaired[0].archetype_fields["qa_expand_layout"] is True


def test_balanced_evidence_cards_use_readable_centered_typography():
    spec = SlideSpec(
        slide_id="s01",
        slide_number=1,
        headline="Guilds Organised Economic Life",
        takeaway="Traders and artisans shared resources and responsibility through guilds",
        bullets=[
            "Guild members shared knowledge about markets, supply, demand, and labour.",
            "Cultivators and artisans set practical rules for their occupations.",
            "Guild autonomy allowed commerce to flourish with limited royal interference.",
        ],
        layout_family=LayoutFamily.TWO_COLUMN,
        archetype_fields={"qa_balanced_cards": True},
    )
    payload = PPTXRenderer.render_presentation([spec], DesignSystem(), deck_title="The Rise of Empires")
    prs = Presentation(io.BytesIO(payload))
    slide = prs.slides[0]
    cards = {shape.name.rsplit("-", 1)[-1]: shape for shape in slide.shapes if shape.name.startswith("evidence-card-")}
    texts = [shape for shape in slide.shapes if shape.name.startswith("evidence-text-")]

    assert len(cards) == len(texts) == 3
    for shape in texts:
        sizes = [run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
        assert sizes and min(sizes) >= 17
        card = cards[shape.name.rsplit("-", 1)[-1]]
        assert shape.top > card.top + Inches(0.25)
        assert shape.top + shape.height < card.top + card.height - Inches(0.25)

    report = PresentationQAAgent.evaluate_presentation([spec], DesignSystem(), payload)
    assert "QA-057" not in {issue.checkpoint_id for issue in report.issues}


def test_repairs_use_stable_action_not_message_text():
    slide = SlideSpec(slide_id="s01", slide_number=1, headline="A title with far too many words for a clean readable presentation slide title")
    report = QAReport(
        issues=[ValidationIssue(
            checkpoint_id="QA-003",
            severity=ValidationSeverity.MEDIUM,
            category=ValidationCategory.CONTENT,
            slide_number=1,
            message="localized message that can change freely",
            repair_action="shorten_title",
        )]
    )
    repaired = RepairAgent.apply_corrections([slide], report)
    assert len(repaired[0].headline.split()) <= 10


def test_image_matcher_preserves_explicit_choices_and_skips_non_image_layouts():
    slides = [
        {"headline": "Solar generation", "layoutHint": "text_image", "imageArtifactId": "solar"},
        {"headline": "Wind generation", "layoutHint": "comparison"},
        {"headline": "Wind generation", "layoutHint": "text_image"},
    ]
    assets = [
        AssetMetadata(asset_id="solar", caption="Solar generation panels"),
        AssetMetadata(asset_id="wind", caption="Wind generation turbines"),
    ]
    ImageMatcher.assign_images_semantically(slides, assets)
    assert slides[0]["imageArtifactId"] == "solar"
    assert "imageArtifactId" not in slides[1]
    assert slides[2]["imageArtifactId"] == "wind"


def test_page_and_figure_citations_do_not_trigger_numeric_chart_check():
    slide = SlideSpec(
        slide_id="s01",
        slide_number=1,
        headline="Evidence across the empire",
        bullets=[
            "Pataliputra appears in Fig. 5.2 (p. 6).",
            "Trade routes are documented in source data.pdf#page=8.",
            "A later monument appears in Figure 5.27 on page 29.",
        ],
        layout_family=LayoutFamily.TWO_COLUMN,
    )

    report = PresentationQAAgent.evaluate_presentation([slide], DesignSystem())
    assert "QA-079" not in {issue.checkpoint_id for issue in report.issues}


def test_relevant_images_are_expected_only_on_image_capable_layouts():
    assets = [AssetMetadata(asset_id=f"empire-{idx}", caption="Ancient empire map and monuments") for idx in range(6)]
    slides = [
        SlideSpec(
            slide_id=f"s{idx + 1:02d}",
            slide_number=idx + 1,
            headline="Ancient empire map and monuments",
            bullets=["Evidence about ancient empire map and monuments."],
            layout_family=LayoutFamily.TEXT_IMAGE if idx < 2 else LayoutFamily.TWO_COLUMN,
            image_artifact_id=f"empire-{idx}" if idx < 2 else None,
            image_caption="Ancient empire map and monuments" if idx < 2 else "",
        )
        for idx in range(6)
    ]

    report = PresentationQAAgent.evaluate_presentation(slides, DesignSystem(), available_assets=assets)
    assert "QA-074" not in {issue.checkpoint_id for issue in report.issues}


def test_two_line_professional_title_is_not_destructively_shortened():
    slide = SlideSpec(
        slide_id="s01",
        slide_number=1,
        headline="An empire is a political unit where a central ruler commands diverse peoples and territories",
        bullets=["A grounded definition."],
        layout_family=LayoutFamily.TWO_COLUMN,
    )

    report = PresentationQAAgent.evaluate_presentation([slide], DesignSystem())
    assert "QA-003" not in {issue.checkpoint_id for issue in report.issues}


def test_number_badges_do_not_create_a_false_body_font_range_failure():
    slide = SlideSpec(
        slide_id="s01",
        slide_number=1,
        headline="What forces drive empire formation?",
        bullets=[
            "Leadership united territories.",
            "Trade generated resources.",
            "Armies secured borders.",
            "Institutions maintained order.",
        ],
        layout_family=LayoutFamily.CARD_GRID,
        layout_hint="numbered_columns",
    )
    payload = PPTXRenderer.render_presentation([slide], DesignSystem(), deck_title="Empires")

    report = PresentationQAAgent.evaluate_presentation([slide], DesignSystem(), pptx_bytes=payload)
    assert "QA-029" not in {issue.checkpoint_id for issue in report.issues}


def test_dynamic_image_allocation_ratio():
    assets = [
        AssetMetadata(asset_id=f"doc-img-{i}", caption=f"Archaeological finding pillar inscription {i}")
        for i in range(1, 6)
    ]
    slides = [
        {"headline": "Ancient Civilizations Overview", "purpose": "Introduction"},
        {"headline": "Archaeological finding pillar inscription 1", "bullets": ["Pillar details."]},
        {"headline": "Administrative Governance", "bullets": ["State apparatus."]},
        {"headline": "Archaeological finding pillar inscription 2", "bullets": ["Territorial borders."]},
        {"headline": "Economic Trade Corridors", "bullets": ["Maritime routes."]},
        {"headline": "Archaeological finding pillar inscription 3", "bullets": ["Coins and weights."]},
        {"headline": "Strategic Inscriptions", "bullets": ["Rock edicts."]},
        {"headline": "Archaeological finding pillar inscription 4", "bullets": ["Edict discovery."]},
        {"headline": "Empire Legacy", "bullets": ["Lasting impact."]},
        {"headline": "Summary & Conclusion", "purpose": "Wrap up"},
    ]
    ImageMatcher.assign_images_semantically(slides, assets)
    assigned_images = [s.get("imageArtifactId") for s in slides if s.get("imageArtifactId")]
    assert len(assigned_images) >= 3
    assert len(assigned_images) == len(set(assigned_images))  # 100% unique


def test_repetitive_box_patterns_detected_and_differentiated():
    slides = [
        SlideSpec(slide_id="s01", slide_number=1, headline="Title", layout_family=LayoutFamily.HERO),
        SlideSpec(slide_id="s02", slide_number=2, headline="Section A", bullets=["Point 1", "Point 2"], layout_family=LayoutFamily.CARD_GRID),
        SlideSpec(slide_id="s03", slide_number=3, headline="Section B", bullets=["Point 3", "Point 4"], layout_family=LayoutFamily.CARD_GRID),
        SlideSpec(slide_id="s04", slide_number=4, headline="Section C", bullets=["Point 5", "Point 6"], layout_family=LayoutFamily.CARD_GRID),
    ]
    report = PresentationQAAgent.evaluate_presentation(slides, DesignSystem())
    failed = {issue.checkpoint_id for issue in report.issues}
    assert "QA-077" in failed  # consecutive_same_layout

    repaired = RepairAgent.apply_corrections(slides, report)
    layouts = [s.layout_family for s in repaired]
    # Check that consecutive identical layouts are broken
    assert layouts[1] != layouts[2] or layouts[2] != layouts[3]


def test_strip_citations_removes_file_and_page_leakage():
    from app.services.design_system import clean_text, strip_citations

    raw_bullets = [
        "Satraps were governors left by overlords to manage far-off territories (source data.pdf#page=15).",
        "Empires maintained armies to conquer and defend borders (source data.pdf#page=6).",
        "Control of rivers and trade networks was a key imperial strategy (source data.pdf#page=8).",
        "Documented evidence in [research_report.pdf#page=22, p. 23] highlights fiscal strength.",
    ]
    cleaned = [clean_text(b) for b in raw_bullets]
    assert cleaned[0] == "Satraps were governors left by overlords to manage far-off territories."
    assert cleaned[1] == "Empires maintained armies to conquer and defend borders."
    assert cleaned[2] == "Control of rivers and trade networks was a key imperial strategy."
    assert "pdf" not in cleaned[3].lower()

    slide = SlideSpec(
        slide_id="s01",
        slide_number=1,
        headline="Post-Mauryan period: regional rulers",
        bullets=raw_bullets,
        layout_family=LayoutFamily.TWO_COLUMN,
    )
    repaired = RepairAgent.apply_corrections([slide], PresentationQAAgent.evaluate_presentation([slide], DesignSystem()))
    for b in repaired[0].bullets:
        assert ".pdf" not in b
        assert "#page=" not in b


def test_explicit_user_prompt_details_and_outlines_preserved():
    from app.agents.requirements_agent import RequirementsAgent
    from app.agents.storyline_agent import StorylineAgent
    from app.services.design_system import fallback_plan

    prompt = """
    Create a presentation on Quantum Computing Architecture:
    Slide 1: Executive Overview - Next-generation computing paradigm
    Slide 2: Physical Realizations & Qubit Hardware [Layout: Two Column]
      - Superconducting Transmon circuits operate near absolute zero
      - Trapped Ion systems offer superior coherence times
    Slide 3: Algorithmic Advantage & Speedup [Layout: Metrics Grid]
      - Shor's algorithm provides exponential speedup for factoring
      - Grover's algorithm provides quadratic speedup for search
    Slide 4: Key Technical Bottlenecks
      - Quantum error correction overhead requires 1000 physical qubits per logical qubit
      - Cryogenic scaling and RF control line density
    Slide 5: Commercial Outlook & 2030 Horizon [Layout: Timeline Band]
      - 2024: Noisy Intermediate-Scale Quantum (NISQ) demonstrations
      - 2027: Early fault-tolerant logical qubit operations
      - 2030: Quantum utility in materials science and cryptography
    """

    # 1. Requirements Agent extracts all 5 slides, custom bullets, and layout hints
    explicit_slides = RequirementsAgent.extract_explicit_slides(prompt)
    assert len(explicit_slides) == 5
    assert explicit_slides[0]["headline"] == "Executive Overview"
    assert "Superconducting" in explicit_slides[1]["bullets"][0]
    assert explicit_slides[2]["layoutHint"] == "metrics_grid"
    assert explicit_slides[4]["layoutHint"] == "timeline_band"

    goal = RequirementsAgent.analyze_requirements(prompt)
    assert goal.target_slide_count == 5
    assert len(goal.explicit_slides) == 5
    assert "EXPLICIT_SLIDE_OUTLINE" in goal.user_directives

    # 2. Storyline Agent generates slide specs directly reflecting user's outline
    specs = StorylineAgent.create_storyline_plan(goal)
    assert len(specs) == 5
    assert specs[0].headline.startswith("Executive Overview")
    assert any("Superconducting" in b for b in specs[1].bullets)
    assert any("2030" in b for b in specs[4].bullets)

    # 3. Fallback Plan directly returns user explicit slides
    fb = fallback_plan(prompt)
    assert len(fb["slides"]) == 5
    assert fb["generationMode"] == "user_explicit_outline"
    assert fb["slides"][0]["headline"] == "Executive Overview"



