import asyncio
import io
import re
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from app.schemas.generation_state import (
    AssetMetadata,
    DesignSystem,
    LayoutFamily,
    PresentationGoal,
    SlideSpec,
    ValidationSeverity,
)
from app.services.archetype_renderer import ArchetypeRenderer
from app.services.renderer import PPTXRenderer
from app.agents.storyline_agent import StorylineAgent
from app.agents.qa_agent import PresentationQAAgent
from app.agents.repair_agent import RepairAgent
from app.services.image_matcher import ImageMatcher, tokenize


def test_no_hardcoded_strings_in_archetype_renderer():
    """Verify that no historical textbook mock strings appear in archetype output."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    spec = SlideSpec(
        slide_number=1,
        headline="Quarterly Cloud Infrastructure Growth",
        key_message="Sustained adoption across enterprise accounts",
        takeaway="Operational reliability remains top priority",
        bullets=[
            "Data Center Expansion: Scaled across 3 major regions",
            "Latency Reduction: Achieved sub-50ms roundtrip delivery",
        ],
        layout_family=LayoutFamily.A24_TIMELINE_BAND,
        archetype_id="A24",
    )
    ds = DesignSystem()

    ArchetypeRenderer.render(slide, "A24", spec, ds)

    all_text = " ".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )

    forbidden_strings = [
        "Swāmi",
        "Peshwa",
        "Amatya",
        "Sachiv",
        "Sumant",
        "Panditrao",
        "Nyayadish",
        "Senapati",
        "Limb",
        "Forts are the core",
        "Universal moral conduct",
        "1630: Foundation & Early Sovereignty",
        "Dharma & Rajyavyavahara",
    ]

    for forbidden in forbidden_strings:
        assert forbidden.lower() not in all_text.lower(), f"Found forbidden hardcoded string: '{forbidden}'"


def test_timeline_chronology_detection():
    """StorylineAgent._determine_layout should route chronological content to LayoutFamily.TIMELINE."""
    raw_slide = {
        "headline": "Major Milestones of Modern Computing",
        "purpose": "Historical development from early mainframes to cloud",
        "bullets": [
            "1945: ENIAC becomes operational as first programmable computer",
            "1969: ARPANET established, creating the foundation for the internet",
            "1981: IBM Personal Computer launched, revolutionizing business computing",
            "1991: World Wide Web made publicly available",
        ],
    }
    goal = PresentationGoal(topic="History of Computing", target_slide_count=5)

    chosen = StorylineAgent._determine_layout(
        index=1,
        total_slides=5,
        hint="",
        recent=[],
        raw_slide=raw_slide,
        assets=[],
        goal=goal,
    )
    assert chosen == LayoutFamily.TIMELINE


def test_a24_timeline_rendering():
    """Verify A24 dynamically renders 4 milestones without Shivaji/Maratha mock data."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    spec = SlideSpec(
        slide_number=2,
        headline="Key Milestones in Renewable Energy Transition",
        bullets=[
            "1997: Kyoto Protocol adopted",
            "2015: Paris Agreement signed by 196 nations",
            "2020: Global solar PV capacity surpasses 700 GW",
            "2025: Renewable generation accounts for 35% of global power",
        ],
        layout_family=LayoutFamily.TIMELINE,
        archetype_id="A24",
    )
    ds = DesignSystem()

    ArchetypeRenderer.render(slide, "A24", spec, ds)

    rendered_text = " ".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )
    assert "1997" in rendered_text
    assert "Paris Agreement" in rendered_text
    assert "Kyoto Protocol" in rendered_text
    assert "1630" not in rendered_text


def test_qa_agent_detects_long_paragraph_title():
    """QA Agent must detect overly long and paragraph-style titles as HIGH severity."""
    specs = [
        SlideSpec(
            slide_number=1,
            headline="Comprehensive Overview of the Strategic Imperatives and Multiple Cross-Functional Initiatives Undertaken During the Entire Fiscal Year.",
            bullets=["Bullet A: Revenue grew by 20%", "Bullet B: Margins improved by 400 bps"],
            layout_family=LayoutFamily.TWO_COLUMN,
        )
    ]
    ds = DesignSystem()
    report = PresentationQAAgent.evaluate_presentation(specs, ds)

    long_title_issues = [
        i for i in report.issues
        if i.checkpoint_id in ("QA-003", "QA-121") or "title" in i.message.lower()
    ]
    assert len(long_title_issues) >= 1
    assert any(i.severity in (ValidationSeverity.HIGH, ValidationSeverity.MEDIUM) for i in long_title_issues)


def test_qa_agent_detects_timeline_chronology_disorder():
    """QA Agent must detect milestones out of chronological sequence."""
    specs = [
        SlideSpec(
            slide_number=2,
            headline="Chronological Evolution of State Policy",
            bullets=[
                "1942: Quit India Movement",
                "1919: Jallianwala Bagh incident",  # Out of order!
                "1947: Indian Independence Act",
            ],
            layout_family=LayoutFamily.TIMELINE,
            archetype_id="A24",
        )
    ]
    ds = DesignSystem()
    report = PresentationQAAgent.evaluate_presentation(specs, ds)

    disorder_issues = [
        i for i in report.issues
        if i.checkpoint_id == "QA-123" or "chronological" in i.message.lower()
    ]
    assert len(disorder_issues) >= 1
    assert disorder_issues[0].severity == ValidationSeverity.HIGH


def test_qa_agent_detects_empty_shape_in_rendered_pptx():
    """QA Agent must inspect rendered shapes and flag empty content boxes."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Add a title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    title_box.text = "Strategic Roadmap"

    # Add an empty content card
    empty_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.5), Inches(4), Inches(2))
    empty_card.name = "content-card-1"
    empty_card.text_frame.text = ""  # Empty!

    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_bytes = pptx_io.getvalue()

    specs = [
        SlideSpec(
            slide_number=1,
            headline="Strategic Roadmap",
            bullets=["Point 1: Strategic Direction"],
            layout_family=LayoutFamily.CARD_GRID,
        )
    ]
    ds = DesignSystem()
    report = PresentationQAAgent.evaluate_presentation(specs, ds, pptx_bytes=pptx_bytes)

    empty_shape_issues = [
        i for i in report.issues
        if i.checkpoint_id == "QA-116" or "empty shape" in i.message.lower()
    ]
    assert len(empty_shape_issues) >= 1


def test_archetype_container_cards_not_flagged_as_empty_shapes():
    """Archetype cards and badges containing nested text must not be flagged as empty shapes."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    spec = SlideSpec(
        slide_number=2,
        slide_id="s02",
        headline="A $XX B addressable market for cloud workflow automation",
        bullets=[
            "TAM for SMB-to-Enterprise SaaS exceeds $XX B",
            "CAGR of 30% for cloud workflow tools",
            "Fragmented landscape: SMB tools lack enterprise security",
        ],
        layout_hint="big_questions",
        archetype_id="A29",
        layout_family=LayoutFamily.CARD_GRID,
    )
    ds = DesignSystem()
    ArchetypeRenderer.render(slide, "A29", spec, ds)

    bio = io.BytesIO()
    prs.save(bio)
    pptx_bytes = bio.getvalue()

    report = PresentationQAAgent.evaluate_presentation([spec], ds, pptx_bytes=pptx_bytes)
    empty_shapes = [i for i in report.issues if i.checkpoint_id == "QA-116"]
    assert len(empty_shapes) == 0, f"Expected 0 empty shape issues on A29 cards, got: {[i.message for i in empty_shapes]}"


def test_repair_agent_rewrites_bad_title():
    """Repair Agent must shorten and clean paragraph titles."""
    bad_spec = SlideSpec(
        slide_number=1,
        headline="The Comprehensive List of Key Milestones and Critical Technical Architecture Changes That Occurred in 2025.",
        bullets=["Milestone A", "Milestone B"],
        layout_family=LayoutFamily.TWO_COLUMN,
    )
    ds = DesignSystem()
    report = PresentationQAAgent.evaluate_presentation([bad_spec], ds)

    repaired = RepairAgent.apply_corrections([bad_spec], report)
    assert len(repaired) == 1
    repaired_title = repaired[0].headline
    assert len(repaired_title.split()) <= 8
    assert not repaired_title.endswith(".")


def test_repair_agent_reorders_timeline():
    """Repair Agent reorder_timeline should sort out-of-order years."""
    disordered_spec = SlideSpec(
        slide_number=1,
        headline="Historical Timeline",
        bullets=[
            "1947: Independence",
            "1857: Revolt",
            "1919: Jallianwala Bagh",
        ],
        layout_family=LayoutFamily.TIMELINE,
        archetype_id="A24",
    )
    ds = DesignSystem()
    report = PresentationQAAgent.evaluate_presentation([disordered_spec], ds)

    repaired = RepairAgent.apply_corrections([disordered_spec], report)
    assert len(repaired) == 1
    bullets = repaired[0].bullets
    assert "1857" in bullets[0]
    assert "1919" in bullets[1]
    assert "1947" in bullets[2]


def test_image_matcher_semantic_allocation():
    """Verify ImageMatcher assigns available valid documentary assets to slides."""
    slides = [
        SlideSpec(slide_number=1, headline="Ancient Mauryan Empire", bullets=["Empire founded by Chandragupta"]),
        SlideSpec(slide_number=2, headline="Ashoka's Edicts and Inscriptions", bullets=["Major Rock Edicts deciphered by James Prinsep"]),
        SlideSpec(slide_number=3, headline="Administrative Structure and Economy", bullets=["Agrarian revenue and coinage system"]),
    ]
    assets = [
        AssetMetadata(
            asset_id="art_ashoka_edicts",
            caption="Ashoka pillar inscription at Sarnath",
            quality_score=0.9,
        ),
        AssetMetadata(
            asset_id="art_mauryan_map",
            caption="Map of the Mauryan Empire territories",
            quality_score=0.9,
        ),
    ]

    ImageMatcher.assign_images_semantically(slides, assets, topic_context="Mauryan Empire")

    assigned_slides = [s for s in slides if s.image_artifact_id]
    assert len(assigned_slides) >= 1
    assert any(s.image_artifact_id == "art_ashoka_edicts" for s in assigned_slides)


def test_repair_agent_handles_diverse_qa_actions():
    """Verify RepairAgent resolves citations, softened claims, bold reduction, and unit inferences."""
    from app.schemas.generation_state import ValidationIssue, ValidationCategory, QAReport

    spec = SlideSpec(
        slide_number=1,
        headline="**Overview**",
        bullets=[
            "Our solution is the only solution and always guarantees 100% uptime (Smith et al., 2020).",
            "**Critical feature with bold formatting**",
        ],
        metrics=[{"value": "85", "label": ""}],
    )

    issues = [
        ValidationIssue(
            checkpoint_id="QA-004",
            severity=ValidationSeverity.MEDIUM,
            category=ValidationCategory.CONTENT,
            slide_number=1,
            slide_id="s01",
            message="Title is too vague",
            repair_action="clarify_title",
        ),
        ValidationIssue(
            checkpoint_id="QA-022",
            severity=ValidationSeverity.MEDIUM,
            category=ValidationCategory.CONTENT,
            slide_number=1,
            slide_id="s01",
            message="Absolute unverified claims",
            repair_action="soften_claim",
        ),
        ValidationIssue(
            checkpoint_id="QA-024",
            severity=ValidationSeverity.MEDIUM,
            category=ValidationCategory.CONTENT,
            slide_number=1,
            slide_id="s01",
            message="Academic citation in body",
            repair_action="move_citation_to_notes",
        ),
        ValidationIssue(
            checkpoint_id="QA-034",
            severity=ValidationSeverity.LOW,
            category=ValidationCategory.DESIGN,
            slide_number=1,
            slide_id="s01",
            message="Overuse of bold",
            repair_action="reduce_bold",
        ),
        ValidationIssue(
            checkpoint_id="QA-038",
            severity=ValidationSeverity.MEDIUM,
            category=ValidationCategory.DATA,
            slide_number=1,
            slide_id="s01",
            message="Metric missing unit or label",
            repair_action="infer_or_flag_units",
        ),
    ]

    report = QAReport(status="needs_repair", checkpoints_passed=115, checkpoints_total=120, issues=issues)
    repaired = RepairAgent.apply_corrections([spec], report, topic="Enterprise Cloud Strategy")

    res = repaired[0]
    # 1. Clarified title
    assert "Enterprise Cloud Strategy" in res.headline or len(res.headline) > len("Overview")
    # 2. Softened claims & removed citation
    assert "Smith et al" not in res.bullets[0]
    assert "Smith et al" in res.speaker_notes
    assert "the only solution" not in res.bullets[0]
    # 3. Reduced bold
    assert "**" not in res.bullets[1]
    # 4. Inferred metric units and label
    assert "%" in res.metrics[0]["value"]
    assert len(res.metrics[0]["label"]) > 0


@pytest.mark.asyncio
async def test_repair_agent_llm_slide_repair(monkeypatch):
    """Verify RepairAgent.repair_slide_with_llm integrates with LLM to repair slide defects."""
    from app.services.provider_router import ProviderRouter
    from app.schemas.generation_state import ValidationIssue, ValidationCategory

    mock_llm_response = {
        "headline": "Streamlined Cloud Workflow Automation",
        "bullets": [
            "Unified control plane reduces reconciliation latency across clusters",
            "Automated policy enforcement mitigates multi-region audit risks",
            "Predictable operational cost scaling replaces bespoke consulting models",
        ],
        "takeaway": "Enterprise-grade governance drives sustainable workflow efficiency.",
        "layoutHint": "two_column",
        "metrics": [{"value": "3x", "label": "Throughput Velocity"}],
    }

    async def fake_call_llm(*args, **kwargs):
        return mock_llm_response

    monkeypatch.setattr(ProviderRouter, "call_llm", fake_call_llm)

    spec = SlideSpec(
        slide_number=2,
        headline="A very bad narrative paragraph title that goes on forever and ever.",
        bullets=["Weak bullet 1", "Weak bullet 2"],
        layout_family=LayoutFamily.CARD_GRID,
    )

    issues = [
        ValidationIssue(
            checkpoint_id="QA-121",
            severity=ValidationSeverity.HIGH,
            category=ValidationCategory.CONTENT,
            slide_number=2,
            slide_id="s02",
            message="Title phrased like narrative sentence",
            repair_action="rewrite_title",
        )
    ]

    repaired_slide = await RepairAgent.repair_slide_with_llm(
        db=None,
        slide=spec,
        issues=issues,
    )

    assert repaired_slide.headline == "Streamlined Cloud Workflow Automation"
    assert len(repaired_slide.bullets) == 3
    assert repaired_slide.takeaway == "Enterprise-grade governance drives sustainable workflow efficiency."
    assert repaired_slide.layout_family == LayoutFamily.TWO_COLUMN
    assert repaired_slide.metrics[0]["value"] == "3x"

