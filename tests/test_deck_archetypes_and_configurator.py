"""Unit tests for consulting deck archetypes and automated PPT design configurator tool."""

import io
import os
import pytest

from pptx import Presentation
from pptx.util import Inches

from app.schemas.generation_state import LayoutFamily, PresentationGoal, SlideSpec
from app.services.deck_archetypes import (
    ArchetypeSelector,
    ColorTokens,
    GRID_SPEC,
    LAYOUT_ARCHETYPES,
    PaletteGenerator,
)
from app.services.design_preset_registry import DesignPresetRegistry
from app.services.renderer import PPTXRenderer
from app.tools.design_auto_configurator import DesignAutoConfigurator


def test_layout_archetypes_catalog_integrity():
    """Verify all 30 archetypes are registered with valid schemas and families."""
    assert len(LAYOUT_ARCHETYPES) == 30
    for arch_id in (f"A{i}" for i in range(1, 31)):
        assert arch_id in LAYOUT_ARCHETYPES
        arch = LAYOUT_ARCHETYPES[arch_id]
        assert "family" in arch
        assert "name" in arch
        assert "when_to_use" in arch
        assert "schema_keys" in arch


def test_new_layout_families_rendering():
    """Verify rendering of newly added layout archetypes (A24-A30)."""
    from app.agents.design_intelligence import DesignIntelligenceAgent

    goal = PresentationGoal(topic="The Rise of the Marathas and Constitutional Federalism", target_slide_count=7)
    ds = DesignIntelligenceAgent.generate_design_system(goal)

    slides = [
        SlideSpec(
            headline="The Big Questions of State Formation",
            eyebrow="CHAPTER 1 · INQUIRY",
            layout_family=LayoutFamily.BIG_QUESTIONS,
            bullets=[
                "Who were the key actors and what drove their mobilization?",
                "What administrative institutions enabled durable statecraft?",
                "How did fiscal and military policies sustain expansion?",
                "What enduring legacy shaped subsequent constitutional design?",
            ],
            takeaway="Four core questions guide the analysis of institutional evolution.",
        ),
        SlideSpec(
            headline="Chronology of the Maratha Sovereign State",
            eyebrow="CHAPTER 2 · TIMELINE",
            layout_family=LayoutFamily.TIMELINE_BAND,
            bullets=[
                "1630: Early Foundations and Clan Roots in the Western Ghats",
                "1657: Naval Power and Coastal Fortress Construction",
                "1674: Coronation as Chhatrapati and Sovereign Recognition",
                "1707: Pan-Indian Expansion and the Peshwa Governance Era",
            ],
            takeaway="A continuous institutional evolution from regional hill autonomy to subcontinental power.",
        ),
        SlideSpec(
            headline="The Ashta Pradhana Mandala",
            eyebrow="CHAPTER 3 · GOVERNANCE",
            layout_family=LayoutFamily.COUNCIL_EIGHT,
            bullets=[
                "Peshwa: Prime Minister overseeing civil and military governance.",
                "Amatya: Finance Minister managing revenue and treasury.",
                "Sachiv: Royal Secretary handling state correspondence.",
                "Mantri: Chronicler keeping official court records.",
                "Senapati: Commander-in-Chief leading armed forces.",
                "Sumant: Foreign Minister managing diplomacy.",
                "Nyayadhish: Chief Justice administering civil justice.",
                "Panditrao: High Priest supervising religious endowments.",
            ],
            takeaway="A structured eight-member council distributing executive authority.",
        ),
        SlideSpec(
            headline="Two Routes to a Federation",
            eyebrow="CHAPTER 4 · COMPARISON",
            layout_family=LayoutFamily.TWO_HIGHWAYS,
            bullets=[
                "Coming Together: Independent sovereign units pool authority to form a larger federal union.",
                "Holding Together: A large unitary nation decides to divide constitutional power between national and regional tiers.",
            ],
            takeaway="Distinct historical pathways produce different balances of regional autonomy.",
        ),
        SlideSpec(
            headline="The Strategic Core of the State: Forts",
            eyebrow="CHAPTER 5 · DOCTRINE",
            layout_family=LayoutFamily.FORTS_QUOTE_EMBLEM,
            bullets=[
                "Military Bastions: Permanent hill fortresses neutralised superior conventional armies.",
                "Fiscal Sanctuaries: Sovereign treasuries and mints remained secure in rugged terrain.",
                "Administrative Hubs: Territorial law and tax administration radiated from fort command.",
            ],
            takeaway="Forts are the core of the state. What is a kingdom without forts? It is like a house without walls.",
        ),
        SlideSpec(
            headline="Defining Federalism and Shared Rule",
            eyebrow="CHAPTER 6 · CONCEPTS",
            layout_family=LayoutFamily.CONCEPT_DEFINITION_IMAGE,
            bullets=[
                "Constitutional Division: Authority divided between central union and constituent states.",
                "Dual Jurisdiction: Each tier possesses guaranteed autonomy in specified legislative fields.",
                "Judicial Guardrails: Supreme Court acts as arbiter of jurisdictional disputes.",
            ],
            takeaway="A system of government where power is constitutionally divided between central and regional tiers.",
        ),
        SlideSpec(
            headline="The Market Value Chain: Cotton to Cloth",
            eyebrow="CHAPTER 7 · VALUE CHAIN",
            layout_family=LayoutFamily.STEPPED_VALUE_CHAIN,
            bullets=[
                "Farmer & Cultivator: Harvests raw cotton pods and sells to regional traders.",
                "Ginning & Spinning Mill: Cleans raw cotton and spins fibers into industrial yarn.",
                "Weaving & Dyeing: Converts yarn into finished textiles and vibrant garments.",
                "Wholesale Merchant: Distributes bulk garments to national distribution centers.",
                "Retail Consumer Outlet: Sells finished garments to shoppers in city markets.",
            ],
            takeaway="Value accumulates across each stage reflecting logistical transport, processing, and margin.",
        ),
    ]

    pptx_bytes = PPTXRenderer.render_presentation(slides, ds, deck_title="Historical & Civic Frameworks")
    assert len(pptx_bytes) > 20000
    report = PPTXRenderer.validate_deck(pptx_bytes, 7)
    assert report["status"] == "passed"
    assert report["overall_quality_score"] >= 75.0


def test_palette_generator_semantic_matching():
    """Verify topic keywords generate expected consulting color tokens."""
    marathas_pal = PaletteGenerator.generate_palette("The Rise of the Maratha Empire")
    assert marathas_pal.ink == "#7E2C22"
    assert marathas_pal.primary == "#E4791F"

    fed_pal = PaletteGenerator.generate_palette("Federalism in the Indian Constitution")
    assert fed_pal.ink == "#0C3B39"
    assert fed_pal.primary == "#0E7C7B"

    econ_pal = PaletteGenerator.generate_palette("Sectors of the Indian Economy")
    assert econ_pal.ink == "#1F3864"
    assert econ_pal.primary == "#2F5597"


def test_archetype_selector_anti_pattern_avoidance():
    """Verify consecutive duplicate archetypes are avoided."""
    first = ArchetypeSelector.select_archetype("comparison")
    second = ArchetypeSelector.select_archetype("comparison", previous_archetype=first)
    assert first != second


def test_archetype_rendering_to_pptx():
    """Verify rendering slides with consulting archetypes produces valid presentation."""
    from app.agents.design_intelligence import DesignIntelligenceAgent
    from app.agents.storyline_agent import StorylineAgent

    goal = PresentationGoal(topic="Federalism Governance Framework", target_slide_count=5)
    specs = StorylineAgent.create_storyline_plan(goal)
    ds = DesignIntelligenceAgent.generate_design_system(goal)

    # Ensure archetypes are assigned
    for s in specs:
        assert s.archetype_id is not None
        assert s.archetype_id.startswith("A")

    pptx_bytes = PPTXRenderer.render_presentation(specs, ds, deck_title="Federalism Governance")
    assert len(pptx_bytes) > 10000

    report = PPTXRenderer.validate_deck(pptx_bytes, 5)
    assert report["status"] == "passed"
    assert report["overall_quality_score"] >= 70.0


def test_design_auto_configurator_extract_and_register():
    """Verify auto-configurator extracts visual language from a PPTX and registers a preset."""
    # Create a minimal synthetic PPTX with custom colors and font
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(5), Inches(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(14, 124, 123)

    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = "Strategic Executive Review"
    p.font.name = "Cambria"
    p.font.size = Inches(0.4)

    buf = io.BytesIO()
    prs.save(buf)
    sample_bytes = buf.getvalue()

    preset = DesignAutoConfigurator.configure_from_pptx(sample_bytes, name="Test Auto Config Preset")
    assert preset["id"].startswith("preset_")
    assert preset["name"] == "Test Auto Config Preset"
    assert "ink" in preset["colors"]
    assert "primary" in preset["colors"]
    assert "title_font" in preset["typography"]
    assert preset["typography"]["title_font"] == "Cambria"

    # Verify preset was registered in registry
    retrieved = DesignPresetRegistry.get_preset(preset["id"])
    assert retrieved is not None
    assert retrieved["id"] == preset["id"]
