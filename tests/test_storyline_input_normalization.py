import io

from pptx import Presentation

from app.agents.design_intelligence import DesignIntelligenceAgent
from app.agents.storyline_agent import StorylineAgent
from app.schemas.generation_state import AssetMetadata, PresentationGoal, SlideSpec
from app.services.image_matcher import ImageMatcher
from app.services.renderer import PPTXRenderer


def test_slide_spec_normalizes_structured_bullets() -> None:
    spec = SlideSpec(
        bullets=[
            {"label": "Nalanda", "value": "427 CE"},
            {"title": "Valabhi", "description": "480 CE"},
            {"text": "Vikramashila flourished from 783–820 CE"},
            1200,
            None,
            {},
        ]
    )

    assert spec.bullets == [
        "Nalanda: 427 CE",
        "Valabhi: 480 CE",
        "Vikramashila flourished from 783–820 CE",
        "1200",
    ]


def test_storyline_accepts_provider_structured_bullets() -> None:
    goal = PresentationGoal(topic="Ancient Indian universities", target_slide_count=1)
    plan = {
        "slides": [
            {
                "purpose": "Compare founding dates",
                "message": "Major centres of learning",
                "bullets": [
                    {"label": "Nalanda", "value": "427 CE"},
                    {"label": "Valabhi", "value": "480 CE"},
                    {"label": "Vikramashila", "value": "783–820 CE"},
                ],
            }
        ]
    }

    slides = StorylineAgent.create_storyline_plan(goal, llm_plan_spec=plan)

    assert slides[0].bullets == [
        "Nalanda: 427 CE",
        "Valabhi: 480 CE",
        "Vikramashila: 783–820 CE",
    ]


def test_image_matcher_normalizes_provider_bullets_before_joining() -> None:
    slides = [
        {
            "headline": "Major centres of learning",
            "purpose": "Compare founding dates",
            "layoutHint": "image_focus",
            "bullets": [{"label": "Nalanda", "value": "427 CE"}],
        }
    ]
    assets = [
        AssetMetadata(
            asset_id="nalanda-image",
            caption="Nalanda major centre of learning",
        )
    ]

    ImageMatcher.assign_images_semantically(slides, assets)

    assert slides[0]["bullets"] == ["Nalanda: 427 CE"]
    assert slides[0]["imageArtifactId"] == "nalanda-image"


def test_image_matcher_rejects_single_weak_keyword_overlap() -> None:
    slides = [
        {
            "headline": "Buddhist relics and religious patronage",
            "purpose": "Explain monuments built for monks",
            "layoutHint": "image_focus",
            "bullets": ["Rulers supported religious communities."],
        }
    ]
    assets = [
        AssetMetadata(
            asset_id="trade-image",
            caption="Rulers controlled rivers and trade networks to secure resources",
        )
    ]

    ImageMatcher.assign_images_semantically(slides, assets)

    assert "imageArtifactId" not in slides[0]


def test_default_section_label_does_not_repeat_the_user_prompt() -> None:
    goal = PresentationGoal(
        topic="Create a professional high-quality 24-slide presentation based only on the attached PDF",
        target_slide_count=1,
    )

    slides = StorylineAgent.create_storyline_plan(goal)

    assert slides[0].section == "Section 1"


def test_numbered_process_preserves_complete_unsplit_sentences() -> None:
    goal = PresentationGoal(topic="Empire growth", target_slide_count=1)
    design = DesignIntelligenceAgent.generate_design_system(goal)
    bullets = [
        "Soldiers marched to battle against neighbouring kingdoms.",
        "Military action addressed external threats.",
        "Defense remained a primary responsibility of rulers.",
    ]
    slide = SlideSpec(
        headline="Military organization",
        bullets=bullets,
        archetype_id="A10",
    )

    payload = PPTXRenderer.render_presentation([slide], design)
    rendered = Presentation(io.BytesIO(payload))
    text = "\n".join(shape.text for shape in rendered.slides[0].shapes if getattr(shape, "has_text_frame", False))

    for bullet in bullets:
        assert bullet in text


def test_numbered_columns_preserve_complete_unsplit_sentences() -> None:
    goal = PresentationGoal(topic="Empire legacy", target_slide_count=1)
    design = DesignIntelligenceAgent.generate_design_system(goal)
    bullets = [
        "Mauryan, Persian, and Greek empires shaped civilizations.",
        "Their influence remains visible in governance, culture, and administration.",
        "Understanding their legacies helps us understand modern society.",
    ]
    slide = SlideSpec(
        headline="Enduring legacies",
        bullets=bullets,
        layout_hint="numbered_columns",
    )

    payload = PPTXRenderer.render_presentation([slide], design)
    rendered = Presentation(io.BytesIO(payload))
    text = "\n".join(shape.text for shape in rendered.slides[0].shapes if getattr(shape, "has_text_frame", False))

    for bullet in bullets:
        assert bullet in text


def test_process_diagram_uses_the_slide_topic_instead_of_product_workflow_copy() -> None:
    goal = PresentationGoal(topic="The rise of empires", target_slide_count=1)
    plan = {
        "slides": [
            {
                "slideId": "s01",
                "headline": "Conquest → Integration → Administration → Legacy",
                "purpose": "Explain the historical sequence",
                "layoutHint": "process_steps",
                "bullets": [
                    "Military conquest established control.",
                    "Institutions integrated diverse territories.",
                    "Administrators governed provinces.",
                    "Successor states inherited imperial practices.",
                ],
            }
        ]
    }

    slide = StorylineAgent.create_storyline_plan(goal, llm_plan_spec=plan)[0]
    labels = [node.label for node in slide.diagram_spec.nodes]
    subtext = [node.subtext for node in slide.diagram_spec.nodes]

    assert labels == ["1. Conquest", "2. Integration", "3. Administration", "4. Legacy"]
    assert subtext == plan["slides"][0]["bullets"]
    assert all("Ingestion" not in label and "Architecture" not in label for label in labels)
