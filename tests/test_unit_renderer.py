"""Unit tests for PPTX renderer."""

import io

from pptx import Presentation

from app.schemas.generation_state import DesignSystem, LayoutFamily, SlideSpec
from app.services.renderer import PPTXRenderer


def test_render_all_slide_layouts():
    spec = {
        "deckTitle": "Comprehensive Strategic Deck",
        "slides": [
            {
                "slideId": "s01",
                "purpose": "Title Hero",
                "message": "Transforming Prompt to Presentation",
                "layoutHint": "hero",
            },
            {
                "slideId": "s02",
                "purpose": "Operational Overview",
                "message": "High-velocity execution across teams",
                "layoutHint": "two_column",
                "bullets": ["Benchmark A hit", "Benchmark B hit"],
            },
            {
                "slideId": "s03",
                "purpose": "KPI Dashboard",
                "message": "Key Performance Metrics",
                "layoutHint": "chart",
            },
            {
                "slideId": "s04",
                "purpose": "Milestone Roadmap",
                "message": "Phased implementation schedule",
                "layoutHint": "timeline",
            },
        ],
    }

    brand = {
        "titleFont": {"name": "Calibri"},
        "colors": {
            "primary": "#1E3A8A",
            "secondary": "#0F172A",
            "accent": "#0D9488",
            "background": "#FFFFFF",
        },
    }

    pptx_bytes = PPTXRenderer.render_deck(spec, brand)
    assert len(pptx_bytes) > 10000
    assert pptx_bytes.startswith(b"PK\x03\x04")

    # Re-open with python-pptx to verify presentation structure
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 4

    # Verify widescreen 16:9 dimensions
    assert int(prs.slide_width.inches) == 13
    assert int(prs.slide_height.inches) == 7


def test_long_title_and_single_point_use_the_available_canvas() -> None:
    headline = "An empire is a political unit where a central ruler commands diverse peoples and territories"
    spec = SlideSpec(
        slide_id="s01",
        slide_number=1,
        headline=headline,
        bullets=["This is one substantive source-grounded conclusion."],
        layout_family=LayoutFamily.DARK_QUOTE,
        dark_background=True,
    )

    payload = PPTXRenderer.render_presentation([spec], DesignSystem(), deck_title="Empires")
    slide = Presentation(io.BytesIO(payload)).slides[0]
    text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
    assert any(shape.text == headline for shape in text_shapes)
    content_cards = [shape for shape in slide.shapes if shape.name == "card"]
    assert content_cards and max(shape.width.inches for shape in content_cards) >= 11.9
