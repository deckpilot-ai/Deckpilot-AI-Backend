"""Tests for multi-agent generation job and PPTX export."""

import asyncio

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.security import hash_password
from app.db.base import Base
from app.models.job import AgentTask, GenerationJob
from app.models.project import Project
from app.models.user import User
from app.services.attachment_pipeline import process_attachment
from app.services.orchestrator import GenerationAlreadyRunningError, JobOrchestrator
from app.services.renderer import PPTXRenderer
from tests.test_projects import create_authenticated_user


def test_generation_pipeline_and_pptx_export(client: TestClient, db_session: Session):
    headers = create_authenticated_user(client, "orchestrator@deckpilot.ai")

    # 1. Create project
    proj_resp = client.post(
        "/api/v1/projects",
        json={"title": "AI Strategy Deck"},
        headers=headers,
    )
    assert proj_resp.status_code == 201
    proj_id = proj_resp.json()["id"]

    # 2. Upload text reference attachment
    upload_resp = client.post(
        f"/api/v1/projects/{proj_id}/attachments",
        files={"file": ("brief.txt", b"deckpilotAI revenue grew by 150% with 98% retention in Q3.", "text/plain")},
        headers=headers,
    )
    assert upload_resp.status_code == 201
    att_data = upload_resp.json()
    assert att_data["status"] == "pending"
    assert att_data["file_name"] == "brief.txt"

    # Extraction runs detached in production; drive it inline for the test.
    asyncio.run(process_attachment(att_data["id"], db=db_session))
    refreshed = client.get(f"/api/v1/projects/{proj_id}/attachments/{att_data['id']}", headers=headers)
    assert refreshed.status_code == 200
    assert refreshed.json()["status"] == "ready"

    # 3. Start generation job
    job_resp = client.post(
        f"/api/v1/projects/{proj_id}/jobs",
        json={"prompt": "Generate a 4-slide executive presentation on AI Strategy", "mode": "generate"},
        headers=headers,
    )
    assert job_resp.status_code == 201
    job_data = job_resp.json()
    assert job_data["status"] == "completed"
    job_id = job_data["job_id"]

    # 4. Inspect job tasks
    get_job_resp = client.get(f"/api/v1/jobs/{job_id}", headers=headers)
    assert get_job_resp.status_code == 200
    job_detail = get_job_resp.json()
    assert len(job_detail["tasks"]) == 8
    for task in job_detail["tasks"]:
        assert task["status"] == "completed"
    assert job_detail["current_step"] == 8
    assert job_detail["total_steps"] == 8
    assert job_detail["progress_percent"] == 100
    assert job_detail["live_message"]
    assert job_detail["progress_events"]
    qa_events = [
        event for event in job_detail["progress_events"]
        if event["agent_type"] == "visual_qa"
    ]
    assert qa_events
    assert any(event.get("phase") == "checking" for event in qa_events)
    assert any(event.get("phase") == "results" for event in qa_events)
    assert any(event.get("phase") == "completed" for event in qa_events)
    assert job_detail["qa_summary"]["checkpoints_total"] == 120
    assert job_detail["qa_summary"]["checkpoints_passed"] >= 1

    # 5. Fetch Deck Version
    deck_resp = client.get(f"/api/v1/projects/{proj_id}/decks/1", headers=headers)
    assert deck_resp.status_code == 200
    deck_data = deck_resp.json()
    assert deck_data["version"] == 1
    assert "slides" in deck_data["spec"]

    # 6. Download PPTX
    download_resp = client.get(f"/api/v1/projects/{proj_id}/decks/1/download", headers=headers)
    assert download_resp.status_code == 200
    assert download_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    # Verify ZIP/PPTX magic bytes (PK\x03\x04)
    assert download_resp.content.startswith(b"PK\x03\x04")


def test_generation_idempotency_key_prevents_duplicate_decks(client: TestClient) -> None:
    headers = create_authenticated_user(client, "idempotent@deckpilot.ai")
    project_id = client.post(
        "/api/v1/projects",
        json={"title": "Idempotent Deck"},
        headers=headers,
    ).json()["id"]
    payload = {
        "prompt": "Create a four-slide operations review",
        "mode": "generate",
        "idempotencyKey": "request-12345678",
    }

    first = client.post(f"/api/v1/projects/{project_id}/jobs", json=payload, headers=headers)
    second = client.post(f"/api/v1/projects/{project_id}/jobs", json=payload, headers=headers)

    assert first.status_code == 201
    assert second.status_code == 201
    assert second.json()["job_id"] == first.json()["job_id"]
    project = client.get(f"/api/v1/projects/{project_id}", headers=headers).json()
    assert project["current_deck_version"] == 1


def test_reexport_preserves_written_slides_and_uses_standalone_cover(client: TestClient, monkeypatch):
    import io

    from PIL import Image, ImageDraw
    from pptx import Presentation

    from app.services.provider_router import ProviderRouter

    headers = create_authenticated_user(client, 'reexport@deckpilot.ai')
    pid = client.post('/api/v1/projects', json={'title': 'History'}, headers=headers).json()['id']
    base = f'/api/v1/projects/{pid}'
    first = client.post(base + '/jobs', json={'prompt': 'Create 4 slides on history'}, headers=headers)
    assert first.status_code == 201
    original = client.get(base + '/decks/1', headers=headers).json()['spec']['slides']
    im = Image.new('RGB', (300, 180), '#bca785')
    ImageDraw.Draw(im).rectangle((40, 30, 240, 160), fill='#4d3928')
    buf = io.BytesIO()
    im.save(buf, format='PNG')
    upload = client.post(base + '/attachments', files={'file': ('cover.png', buf.getvalue(), 'image/png')}, headers=headers)
    assert upload.status_code == 201

    async def no_rewriting(**kwargs):
        assert kwargs['agent_type'] not in ('deck_planner', 'slide_writer', 'font_brand_detection')
        return {'text': 'Re-exported the existing slides.'}

    monkeypatch.setattr(ProviderRouter, 'call_llm', no_rewriting)
    result = client.post(base + '/jobs', json={'prompt': 'Re-export the existing deck', 'mode': 'export'}, headers=headers)
    assert result.status_code == 201 and result.json()['status'] == 'completed'
    updated = client.get(base + '/decks/2', headers=headers).json()['spec']['slides']
    assert [(s.get('message'), s.get('bullets')) for s in original] == [(s.get('message'), s.get('bullets')) for s in updated]
    assert updated[0]['imageArtifactId']
    payload = client.get(base + '/decks/2/download', headers=headers).content
    deck = Presentation(io.BytesIO(payload))
    assert any(s.shape_type == 13 for s in deck.slides[0].shapes)


def test_project_allows_only_one_active_generation(client: TestClient, db_session: Session) -> None:
    headers = create_authenticated_user(client, "active-slot@deckpilot.ai")
    project_id = client.post(
        "/api/v1/projects",
        json={"title": "Active Job Slot"},
        headers=headers,
    ).json()["id"]

    JobOrchestrator.create_job(
        db_session,
        project_id=project_id,
        user_id="ignored-by-orchestrator",
        idempotency_key="active-job-one",
    )

    with pytest.raises(GenerationAlreadyRunningError):
        JobOrchestrator.create_job(
            db_session,
            project_id=project_id,
            user_id="ignored-by-orchestrator",
            idempotency_key="active-job-two",
        )

    blocked_delete = client.delete(f"/api/v1/projects/{project_id}", headers=headers)
    assert blocked_delete.status_code == 409


@pytest.mark.asyncio
async def test_generation_failure_releases_active_slot_and_fails_remaining_tasks(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    isolated_engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(isolated_engine)
    db_session = Session(isolated_engine)
    user = User(email="failed-job@example.com", password_hash=hash_password("password12345"))
    db_session.add(user)
    db_session.flush()
    project = Project(user_id=user.id, title="Failure Recovery")
    db_session.add(project)
    db_session.commit()
    job, _created = JobOrchestrator.create_job(
        db_session,
        project_id=project.id,
        user_id=user.id,
        idempotency_key="failed-job-request",
    )
    job_id = job.id

    def fail_renderer(*_args, **_kwargs):
        raise RuntimeError("simulated renderer failure")

    monkeypatch.setattr(PPTXRenderer, "render_deck", fail_renderer)
    with pytest.raises(RuntimeError, match="simulated renderer failure"):
        await JobOrchestrator.run_job(
            db_session,
            job_id,
            "Create a failure recovery deck",
            user.id,
        )

    failed_job = db_session.get(GenerationJob, job_id)
    assert failed_job is not None
    assert failed_job.status == "permanently_failed"
    assert failed_job.active_slot is None
    unfinished = db_session.scalars(
        select(AgentTask).where(
            AgentTask.job_id == job_id,
            AgentTask.status.in_(("pending", "running")),
        )
    ).all()
    assert unfinished == []
    db_session.close()
    isolated_engine.dispose()
