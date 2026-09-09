from pptx import Presentation

prs = Presentation("Expected PPT/Accelerating Growth Our Journey into Enterprise SaaS_v1.pptx")
print(f"Accelerating Growth SaaS deck: {len(prs.slides)} slides")
for idx, slide in enumerate(prs.slides):
    texts = [s.text.strip().replace('\n', ' ') for s in slide.shapes if s.has_text_frame and s.text.strip()]
    charts = [s for s in slide.shapes if s.has_chart]
    tables = [s for s in slide.shapes if s.has_table]
    images = [s for s in slide.shapes if hasattr(s, 'image')]
    print(f"Slide {idx+1}: shapes={len(slide.shapes)}, charts={len(charts)}, tables={len(tables)}, imgs={len(images)}")
    for t in texts[:4]:
        print(f"   text: {t[:100]}")
    if charts:
        for c in charts:
            print(f"   chart_type={c.chart.chart_type}, title={c.chart.has_title and c.chart.chart_title.text_frame.text}")
