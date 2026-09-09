import glob
from pptx import Presentation

for path in glob.glob("Expected PPT/*.pptx"):
    print(f"\n==================================================")
    print(f"FILE: {path}")
    prs = Presentation(path)
    for idx, slide in enumerate(prs.slides):
        texts = [s.text.strip().replace('\n', ' ') for s in slide.shapes if s.has_text_frame and s.text.strip()]
        charts = [str(s.chart.chart_type) for s in slide.shapes if s.has_chart]
        tables = [f"{len(s.table.rows)}x{len(s.table.columns)}" for s in slide.shapes if s.has_table]
        images = [f"{s.image.size}" for s in slide.shapes if hasattr(s, 'image')]
        print(f"Slide {idx+1}: {len(slide.shapes)} shapes | charts={charts} | tables={tables} | imgs={images}")
        for t in texts[:4]:
            print(f"    txt: {t[:90]}")
