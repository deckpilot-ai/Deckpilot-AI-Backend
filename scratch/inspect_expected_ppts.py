import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os

files = [
    "c:/DeckPilotAI/Expected PPT/Federalism (2).pptx",
    "c:/DeckPilotAI/Expected PPT/The Rise of the Marathas (4).pptx",
    "c:/DeckPilotAI/Expected PPT/The Rise of Empires (3).pptx"
]

for path in files:
    if not os.path.exists(path):
        continue
    prs = pptx.Presentation(path)
    print("=" * 60)
    print(f"PPT: {os.path.basename(path)}")
    print(f"Dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f}")
    print(f"Total Slides: {len(prs.slides)}")
    
    for i, slide in enumerate(prs.slides):
        if i >= 4:
            break
        print(f"\n  --- Slide {i+1} ---")
        shapes_summary = []
        for s in slide.shapes:
            info = f"[{s.shape_type}] '{s.name}' pos=({s.left.inches:.2f}, {s.top.inches:.2f}, {s.width.inches:.2f}, {s.height.inches:.2f})"
            if s.has_text_frame:
                txt = s.text.replace('\n', ' ')[:50]
                fonts = set()
                sizes = set()
                colors = set()
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.name: fonts.add(r.font.name)
                        if r.font.size: sizes.add(round(r.font.size.pt, 1))
                        try:
                            if r.font.color and r.font.color.rgb:
                                colors.add(str(r.font.color.rgb))
                        except Exception:
                            pass
                info += f" | txt='{txt}' | fonts={list(fonts)[:2]} sizes={list(sizes)[:2]} colors={list(colors)[:2]}"
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                info += " [IMAGE]"
            if s.shape_type == MSO_SHAPE_TYPE.TABLE:
                info += f" [TABLE {len(s.table.rows)}x{len(s.table.columns)}]"
            shapes_summary.append(info)
        for sinfo in shapes_summary[:8]:
            print("   ", sinfo)
        if len(shapes_summary) > 8:
            print(f"    ... +{len(shapes_summary) - 8} more shapes")
